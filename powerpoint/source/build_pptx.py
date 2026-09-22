"""Build the SAUS PowerPoint template and demo deck from the brand assets.

    python build_pptx.py

Writes ../SAUS-template.potx (slide master and layouts, no slides) and
../SAUS-slides-demo.pptx (a tour of every layout, mirroring the Beamer demo).
The design follows latex/beamerthemeSAUS.sty: every measurement below is the
Beamer value (16 cm x 9 cm page) scaled to PowerPoint's 33.867 cm x 19.05 cm.

Needs python-pptx (pip install python-pptx).
"""
import os
import zipfile
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.parts.slide import SlideLayoutPart
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
ASSETS = os.path.join(ROOT, "latex", "assets")
ART = os.path.join(HERE, "art")
OUT = os.path.normpath(os.path.join(HERE, ".."))

# ---- Brand -----------------------------------------------------------------
MAROON, MAROON_DARK, MAROON_LIGHT = "6B2A4C", "4A1C34", "8E4166"
BLUE, BLUE_DARK, GOLD = "1A1A91", "131368", "C39B3E"
MIST, INK, SLATE, RULE = "F4EDF1", "231C21", "6A626B", "D9C9D2"
GREEN, RED, WHITE = "2E6B4F", "B3261E", "FFFFFF"
FOOT_TITLE = "C4A9B7"            # sausmist!65!sausmaroon, as in the Beamer footer
BLUE_TINT = "F1F1F8"             # sausblue!6
HEAD, BODY, LIGHT, LABEL, SERIF = ("Fira Sans SemiBold", "Fira Sans",
                                   "Fira Sans Light", "Fira Sans Medium", "EB Garamond")

UNIVERSITY = "The Shaikh Ayaz University, Shikarpur"
MOTTO = "Revival of Educational Glory of Shikarpur"
CONTACT = "saus.edu.pk  ·  Main Road, Shikarpur, Sindh  ·  +92 726 512039"

LOGO = os.path.join(ASSETS, "saus-logo.png")
LOGO_WHITE = os.path.join(ASSETS, "saus-logo-white.png")
LOGO_RATIO = 384 / 484                      # crest width / height
TITLE_PHOTO = os.path.join(ASSETS, "saus-title-bg.jpg")
SINDHI_WHITE = os.path.join(ART, "sindhi-white.png")
SINDHI_MAROON = os.path.join(ART, "sindhi-maroon.png")
URDU_MAROON = os.path.join(ART, "urdu-maroon.png")
SINDHI_RATIO, URDU_RATIO = 3258 / 437, 2861 / 746

# ---- Geometry (cm) -----------------------------------------------------------
W, H = 33.867, 19.05
M = 2.33                        # side margin (1.1 cm in Beamer)
BAND = 0.72                     # footer band
BAND_Y = H - BAND
RULE_Y = 3.2                    # title rule under the frame title
BODY_Y = 3.75
BODY_B = BAND_Y - 0.6
TITLE_RIGHT = W - 6.03          # frame title stops short of the corner crest
CREST_H = 2.43

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')


def emu(cm):
    return int(round(cm * 360000))


def xfrm(x, y, w, h):
    return (f'<a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/>'
            f'<a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>')


def fill(color, alpha=None):
    if color is None:
        return "<a:noFill/>"
    a = f'<a:alpha val="{int(alpha * 1000)}"/>' if alpha is not None else ""
    return f'<a:solidFill><a:srgbClr val="{color}">{a}</a:srgbClr></a:solidFill>'


def rpr(size=None, color=None, font=None, bold=None, italic=None, spc=None, caps=False,
        tag="a:rPr"):
    attrs = ['lang="en-GB"'] if tag == "a:rPr" else []
    if size:
        attrs.append(f'sz="{int(size * 100)}"')
    if bold is not None:
        attrs.append(f'b="{int(bold)}"')
    if italic is not None:
        attrs.append(f'i="{int(italic)}"')
    if spc:
        attrs.append(f'spc="{int(spc)}"')
    if caps:
        attrs.append('cap="all"')
    inner = fill(color) if color else ""
    if font:
        inner += f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/>'
    return f'<{tag} {" ".join(attrs)}>{inner}</{tag}>'


class Shapes:
    """Appends shapes to one part's spTree, numbering them."""

    def __init__(self, part, sptree):
        self.part, self.tree = part, sptree
        self.n = max([int(e.get("id")) for e in sptree.iter(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")] + [1])

    def _id(self):
        self.n += 1
        return self.n

    def add(self, xml):
        self.tree.append(parse_xml(xml))

    def rect(self, x, y, w, h, color, alpha=None, name="Rectangle", geom="rect"):
        self.add(f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self._id()}" name="{name}"/>'
                 f'<p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}'
                 f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>{fill(color, alpha)}'
                 f'<a:ln><a:noFill/></a:ln></p:spPr></p:sp>')

    def fade(self, x, y, w, h, color, alpha, name="Fade"):
        """Colour fading from alpha (left) to clear (right), like TikZ fading=east."""
        self.add(f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self._id()}" name="{name}"/>'
                 f'<p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}'
                 f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:gradFill rotWithShape="1">'
                 f'<a:gsLst><a:gs pos="0"><a:srgbClr val="{color}"><a:alpha val="{int(alpha * 1000)}"/>'
                 f'</a:srgbClr></a:gs><a:gs pos="100000"><a:srgbClr val="{color}"><a:alpha val="0"/>'
                 f'</a:srgbClr></a:gs></a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill>'
                 f'<a:ln><a:noFill/></a:ln></p:spPr></p:sp>')

    def brand_rule(self, x, y, w, scale=1.0):
        """The signature rule: maroon, gold, then a hairline to width w."""
        t, m, g = 0.164 * scale, 1.2 * scale, 0.52 * scale
        self.rect(x, y, m, t, MAROON, name="Rule maroon")
        self.rect(x + m, y, g, t, GOLD, name="Rule gold")
        hl = 0.045 * scale
        self.rect(x + m + g, y + (t - hl) / 2, w - m - g, hl, RULE, name="Rule hairline")

    def text(self, x, y, w, h, paras, anchor="t", name="Text", wrap=True):
        """paras: list of (runs, align, space_before_pt); runs: list of (text, rPr-kwargs)."""
        body = ""
        for runs, algn, before in paras:
            ppr = f'<a:pPr algn="{algn}"><a:spcBef><a:spcPts val="{int(before * 100)}"/></a:spcBef></a:pPr>'
            rs = "".join(f"<a:r>{rpr(**fmt)}<a:t>{escape(t)}</a:t></a:r>" for t, fmt in runs)
            body += f"<a:p>{ppr}{rs}</a:p>"
        self.add(f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self._id()}" name="{name}"/>'
                 f'<p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}'
                 f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
                 f'<p:txBody><a:bodyPr wrap="{"square" if wrap else "none"}" lIns="0" tIns="0" rIns="0" '
                 f'bIns="0" anchor="{anchor}"><a:noAutofit/></a:bodyPr><a:lstStyle/>{body}</p:txBody></p:sp>')

    def pic(self, path, x, y, w, h, alpha=None, name="Picture", clip=True):
        """Picture at (x, y, w, h), cropped to the slide so nothing hangs off the edge."""
        _, rid = self.part.get_or_add_image_part(path)
        l, t, r, b = 0, 0, 0, 0
        if clip:
            l, t = max(0, -x) / w, max(0, -y) / h
            r, b = max(0, x + w - W) / w, max(0, y + h - H) / h
        crop = ""
        if any((l, t, r, b)):
            crop = (f'<a:srcRect l="{int(l * 1e5)}" t="{int(t * 1e5)}" '
                    f'r="{int(r * 1e5)}" b="{int(b * 1e5)}"/>')
            x, y = x + l * w, y + t * h
            w, h = w * (1 - l - r), h * (1 - t - b)
        amt = f'<a:alphaModFix amt="{int(alpha * 1000)}"/>' if alpha is not None else ""
        self.add(f'<p:pic {NS}><p:nvPicPr><p:cNvPr id="{self._id()}" name="{name}"/>'
                 f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/>'
                 f'</p:nvPicPr><p:blipFill><a:blip r:embed="{rid}">{amt}</a:blip>{crop}'
                 f'<a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr>{xfrm(x, y, w, h)}'
                 f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')

    def ph(self, kind, idx, x, y, w, h, prompt="", size=None, color=None, font=None,
           bold=None, italic=None, spc=None, caps=False, algn="l", anchor="t",
           bullets=False, box=None, alpha=None, ins=(0, 0, 0, 0), lnspc=None, name=None):
        """A placeholder. kind: title, ctrTitle, subTitle, body, obj, pic, ftr, sldNum."""
        attrs = []
        if kind != "obj":
            attrs.append(f'type="{kind}"')
        if idx is not None:
            attrs.append(f'idx="{idx}"')
        if kind in ("body", "obj", "pic", "subTitle"):
            attrs.append('hasCustomPrompt="1"' if prompt else "")
        lvl1 = ""
        if not bullets:
            ls = f'<a:lnSpc><a:spcPct val="{int(lnspc * 1000)}"/></a:lnSpc>' if lnspc else ""
            lvl1 = (f'<a:lvl1pPr marL="0" indent="0" algn="{algn}">{ls}<a:spcBef><a:spcPts val="0"/>'
                    f'</a:spcBef><a:buNone/>'
                    f'{rpr(size, color, font, bold, italic, spc, caps, tag="a:defRPr")}</a:lvl1pPr>')
        l, t, r, b = (emu(v) for v in ins)
        body_pr = (f'<a:bodyPr lIns="{l}" tIns="{t}" rIns="{r}" bIns="{b}" anchor="{anchor}">'
                   f'<a:normAutofit/></a:bodyPr>')
        if kind == "sldNum":
            para = ('<a:p><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
                    '<a:rPr lang="en-GB"/><a:t>‹#›</a:t></a:fld></a:p>')
        else:
            para = (f'<a:p><a:r><a:rPr lang="en-GB"/><a:t>{escape(prompt)}</a:t></a:r></a:p>'
                    if prompt else '<a:p><a:endParaRPr lang="en-GB"/></a:p>')
        geom = '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>' + (fill(box, alpha) if box else "")
        nm = name or {"title": "Title", "ctrTitle": "Title", "subTitle": "Subtitle",
                      "pic": "Picture Placeholder", "ftr": "Footer", "sldNum": "Slide Number"
                      }.get(kind, "Text Placeholder")
        tx = "" if kind == "pic" else f'<p:txBody>{body_pr}<a:lstStyle>{lvl1}</a:lstStyle>{para}</p:txBody>'
        self.add(f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self._id()}" name="{nm} {self.n}"/>'
                 f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph {" ".join(attrs)}/>'
                 f'</p:nvPr></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}{geom}</p:spPr>{tx}</p:sp>')


# ---- Theme -------------------------------------------------------------------
def set_theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    from lxml import etree
    root = etree.fromstring(part.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    root.set("name", "SAUS")
    clr = root.find(f".//{a}clrScheme")
    clr.set("name", "SAUS")
    scheme = {"dk1": INK, "lt1": WHITE, "dk2": MAROON_DARK, "lt2": MIST,
              "accent1": MAROON, "accent2": BLUE, "accent3": GOLD, "accent4": GREEN,
              "accent5": RED, "accent6": SLATE, "hlink": BLUE, "folHlink": MAROON_LIGHT}
    for key, val in scheme.items():
        el = clr.find(f"{a}{key}")
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{a}srgbClr").set("val", val)
    fonts = root.find(f".//{a}fontScheme")
    fonts.set("name", "SAUS")
    for key, face in (("majorFont", HEAD), ("minorFont", BODY)):
        f = fonts.find(f"{a}{key}")
        f.find(f"{a}latin").set("typeface", face)
        f.find(f"{a}cs").set("typeface", "Lateef")
        for extra in f.findall(f"{a}font"):     # per-script overrides from the default theme
            if extra.get("script") == "Arab":
                extra.set("typeface", "Lateef")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---- Master --------------------------------------------------------------------
def text_styles():
    def lvl(n, marl, indent, size, bu_color, bu_pct, before, char="§", font="Wingdings",
            color=INK):
        return (f'<a:lvl{n}pPr marL="{emu(marl)}" indent="{emu(indent)}" algn="l">'
                f'<a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="{before * 100}"/>'
                f'</a:spcBef><a:buClr><a:srgbClr val="{bu_color}"/></a:buClr>'
                f'<a:buSzPct val="{bu_pct * 1000}"/><a:buFont typeface="{font}" pitchFamily="2" '
                f'charset="2"/><a:buChar char="{char}"/><a:defRPr sz="{size * 100}" kern="1200">'
                f'{fill(color)}<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
                f'<a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{n}pPr>')
    body = (lvl(1, 0.75, -0.75, 20, MAROON, 70, 10) + lvl(2, 1.5, -0.6, 18, BLUE, 60, 4)
            + lvl(3, 2.2, -0.55, 16, GOLD, 55, 3))
    for n in range(4, 10):
        body += lvl(n, 2.2 + 0.6 * (n - 3), -0.5, 16, SLATE, 90, 2, "–", "Arial", SLATE)
    title = (f'<a:lvl1pPr algn="l"><a:lnSpc><a:spcPct val="95000"/></a:lnSpc><a:spcBef>'
             f'<a:spcPct val="0"/></a:spcBef><a:buNone/><a:defRPr sz="2800" kern="1200">'
             f'{fill(MAROON)}<a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/>'
             f'<a:cs typeface="+mj-cs"/></a:defRPr></a:lvl1pPr>')
    other = (f'<a:defPPr><a:defRPr lang="en-GB"/></a:defPPr><a:lvl1pPr><a:defRPr sz="1800">'
             f'{fill(INK)}<a:latin typeface="+mn-lt"/><a:cs typeface="+mn-cs"/></a:defRPr></a:lvl1pPr>')
    return parse_xml(f'<p:txStyles {NS}><p:titleStyle>{title}</p:titleStyle>'
                     f'<p:bodyStyle>{body}</p:bodyStyle><p:otherStyle>{other}</p:otherStyle></p:txStyles>')


def footer(s, dark_label=True):
    """Footer band: progress-style rule, university name, footer text, slide number."""
    s.rect(0, BAND_Y - 0.1, W, 0.1, MAROON_LIGHT, name="Footer rule")
    s.rect(0, BAND_Y, W, BAND, MAROON, name="Footer band")
    s.text(M, BAND_Y, 12, BAND, [([(UNIVERSITY, dict(size=9, color=WHITE, font=LABEL, spc=120,
                                                    caps=True))], "l", 0)],
           anchor="ctr", name="Footer university")
    s.ph("ftr", 11, W / 2 - 7, BAND_Y, 14, BAND, size=9, color=FOOT_TITLE, algn="ctr",
         anchor="ctr")
    s.ph("sldNum", 12, W - M - 3, BAND_Y, 3, BAND, size=9, color=WHITE, algn="r", anchor="ctr")


def frame_furniture(s):
    """Corner crest and the rule under the frame title."""
    s.pic(LOGO, W - M - CREST_H * LOGO_RATIO, 0.89, CREST_H * LOGO_RATIO, CREST_H, name="Crest")
    s.brand_rule(M, RULE_Y, TITLE_RIGHT - M)


def build_master(prs):
    master = prs.slide_master
    el = master._element
    tree = el.cSld.spTree
    for shape in list(tree)[2:]:
        tree.remove(shape)
    csld = el.cSld
    for bg in csld.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}bg"):
        csld.remove(bg)
    csld.insert(0, parse_xml(f'<p:bg {NS}><p:bgPr>{fill(WHITE)}<a:effectLst/></p:bgPr></p:bg>'))
    s = Shapes(master.part, tree)
    s.ph("title", None, M, 0.6, TITLE_RIGHT - M, RULE_Y - 0.25 - 0.6, "Click to edit title",
         anchor="b", bullets=True)
    s.ph("body", 1, M, BODY_Y, W - 2 * M, BODY_B - BODY_Y, "Click to edit text", bullets=True)
    frame_furniture(s)
    footer(s)
    old = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles")
    el.replace(old, text_styles())


# ---- Layouts -------------------------------------------------------------------
LAYOUT_ID = [2147483649]


def new_layout(prs, name, kind=None, show_master=True, bg=None):
    package = prs.part.package
    partname = package.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
    bg_xml = f"<p:bg><p:bgPr>{fill(bg)}<a:effectLst/></p:bgPr></p:bg>" if bg else ""
    type_attr = f' type="{kind}"' if kind else ""
    show = "" if show_master else ' showMasterSp="0"'
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<p:sldLayout {NS}{type_attr} preserve="1"{show}><p:cSld name="{name}">{bg_xml}'
           f'<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
           f'<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/>'
           f'<a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld>'
           f'<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>')
    part = SlideLayoutPart.load(partname, CT.PML_SLIDE_LAYOUT, package, xml.encode("utf-8"))
    part.relate_to(prs.slide_master.part, RT.SLIDE_MASTER)
    rid = prs.slide_master.part.relate_to(part, RT.SLIDE_LAYOUT)
    lst = prs.slide_master._element.get_or_add_sldLayoutIdLst()
    lst.append(parse_xml(f'<p:sldLayoutId {NS} id="{LAYOUT_ID[0]}" r:id="{rid}"/>'))
    LAYOUT_ID[0] += 1
    return Shapes(part, part._element.cSld.spTree)


def drop_default_layouts(prs):
    master = prs.slide_master
    lst = master._element.get_or_add_sldLayoutIdLst()
    rels = master.part.rels
    for entry in list(lst):
        rid = entry.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        lst.remove(entry)
        rels.pop(rid)


def title_standard(s, top=0.6):
    s.ph("title", None, M, top, TITLE_RIGHT - M, RULE_Y - 0.25 - top, "Click to add title",
         anchor="b", bullets=True)


def title_with_subtitle(s):
    s.ph("title", None, M, 0.35, TITLE_RIGHT - M, 1.95, "Click to add title", anchor="b",
         bullets=True)
    s.ph("body", 10, M, 2.33, TITLE_RIGHT - M, 0.72, "Subtitle", size=16, color=SLATE,
         font=LIGHT, name="Subtitle")


def wordmark_dark(s):
    """White crest with the university name, motto and Sindhi name (title slides)."""
    ch = 3.18
    s.pic(LOGO_WHITE, M, 1.69, ch * LOGO_RATIO, ch, name="Crest")
    x = M + ch * LOGO_RATIO + 0.95
    s.text(x, 1.95, 20, 1.45, [
        ([(UNIVERSITY, dict(size=13, color=WHITE, font=LABEL, spc=150, caps=True))], "l", 0),
        ([(MOTTO, dict(size=14, color=GOLD, font=SERIF, italic=True))], "l", 3)],
        name="University name")
    sh = 0.8
    s.pic(SINDHI_WHITE, x, 3.6, sh * SINDHI_RATIO, sh, name="Sindhi name")


def title_meta(s):
    """Department, presenter, occasion and the contact line along the bottom."""
    s.ph("body", 13, M, 12.75, 0.8 * W, 0.62, "Department of ...", size=12, color=GOLD,
         font=LABEL, spc=150, caps=True, anchor="b", name="Department")
    s.ph("body", 14, M, 13.45, 0.8 * W, 0.85, "Presenter name", size=18, color=WHITE,
         anchor="ctr", name="Presenter")
    s.ph("body", 15, M, 14.35, 0.8 * W, 0.75, "Occasion  ·  Date", size=14,
         color=MIST, name="Occasion and date")
    s.rect(M, H - 2.16, 0.86 * W, 0.06, GOLD, name="Gold rule")
    s.text(M, H - 1.62, 10, 0.6, [([("saus.edu.pk", dict(size=10, color=MIST))], "l", 0)],
           anchor="b", name="Website")
    s.text(W - M - 16, H - 1.62, 16, 0.6, [([("Main Road, Shikarpur, Sindh  ·  +92 726 512039",
                                             dict(size=10, color=MIST))], "r", 0)],
           anchor="b", name="Address")


def title_headline(s):
    s.ph("ctrTitle", None, M, 4.9, 0.78 * W, 4.75, "Presentation title", size=36,
         color=WHITE, font=HEAD, anchor="b", lnspc=95)
    s.rect(M, 9.9, 5.5, 0.164, GOLD, name="Gold rule")
    s.ph("subTitle", 1, M, 10.25, 0.78 * W, 1.6, "Subtitle", size=20, color=MIST,
         font=LIGHT)


def layout_title_photo(prs):
    s = new_layout(prs, "Title Slide (photo)", "title", show_master=False, bg=MAROON_DARK)
    s.pic(TITLE_PHOTO, 0, 0, W, H, name="Campus photograph")
    s.rect(0, 0, W, H, MAROON_DARK, alpha=84, name="Maroon wash")
    s.fade(0, 0, W, H, MAROON_DARK, 55)
    wordmark_dark(s)
    title_headline(s)
    title_meta(s)


def layout_title_solid(prs):
    s = new_layout(prs, "Title Slide (solid)", "title", show_master=False, bg=MAROON_DARK)
    ch = 1.18 * H
    s.pic(LOGO_WHITE, W + 2.33 - ch * LOGO_RATIO, (H - ch) / 2, ch * LOGO_RATIO, ch, alpha=10,
          name="Crest watermark")
    s.fade(0, 0, W, H, MAROON, 40)
    wordmark_dark(s)
    title_headline(s)
    title_meta(s)


def layout_title_split(prs):
    s = new_layout(prs, "Title Slide (split)", "title", show_master=False, bg=MIST)
    band = 0.36 * H
    s.rect(0, H - band, W, band, MAROON_DARK, name="Band")
    s.rect(0, H - band - 0.14, W, 0.14, GOLD, name="Gold edge")
    ch = 2.43
    s.pic(LOGO, (W - ch * LOGO_RATIO) / 2, 0.85, ch * LOGO_RATIO, ch, name="Crest")
    s.text(M, 3.45, W - 2 * M, 1.25, [
        ([(UNIVERSITY, dict(size=12, color=MAROON, font=LABEL, spc=150, caps=True))], "ctr", 0),
        ([(MOTTO, dict(size=13, color=SLATE, font=SERIF, italic=True))], "ctr", 2)],
        name="University name")
    sh = 0.66
    s.pic(SINDHI_MAROON, (W - sh * SINDHI_RATIO) / 2, 4.8, sh * SINDHI_RATIO, sh,
          name="Sindhi name")
    s.ph("ctrTitle", None, M, 5.65, W - 2 * M, 3.15, "Presentation title", size=36,
         color=MAROON, font=HEAD, algn="ctr", anchor="ctr", lnspc=95)
    s.ph("subTitle", 1, M, 8.8, W - 2 * M, 1.3, "Subtitle", size=20, color=SLATE, font=LIGHT,
         algn="ctr")
    top = H - band
    s.ph("body", 13, M, top + 0.75, W - 2 * M, 0.6, "Department of ...", size=12, color=GOLD,
         font=LABEL, spc=150, caps=True, algn="ctr", anchor="b", name="Department")
    s.ph("body", 14, M, top + 1.45, W - 2 * M, 0.85, "Presenter name", size=18, color=WHITE,
         algn="ctr", anchor="ctr", name="Presenter")
    s.ph("body", 15, M, top + 2.35, W - 2 * M, 0.75, "Occasion  ·  Date", size=14,
         color=MIST, algn="ctr", name="Occasion and date")
    s.text(M, H - 1.55, W - 2 * M, 0.55, [([(CONTACT, dict(size=10, color=MIST))], "ctr", 0)],
           anchor="b", name="Contact")


def layout_content(prs):
    s = new_layout(prs, "Title and Content", "obj")
    title_standard(s)
    s.ph("obj", 1, M, BODY_Y, W - 2 * M, BODY_B - BODY_Y, "Click to add text", bullets=True)
    footer_placeholders(s)


def layout_content_subtitle(prs):
    s = new_layout(prs, "Title, Subtitle and Content")
    title_with_subtitle(s)
    s.ph("obj", 1, M, BODY_Y, W - 2 * M, BODY_B - BODY_Y, "Click to add text", bullets=True)
    footer_placeholders(s)


def layout_two(prs):
    s = new_layout(prs, "Two Content", "twoObj")
    title_standard(s)
    cw = (W - 2 * M - 1.2) / 2
    s.ph("obj", 1, M, BODY_Y, cw, BODY_B - BODY_Y, "Click to add text", bullets=True)
    s.ph("obj", 2, M + cw + 1.2, BODY_Y, cw, BODY_B - BODY_Y, "Click to add text", bullets=True)
    footer_placeholders(s)


def layout_blocks(prs):
    s = new_layout(prs, "Comparison (blocks)", "twoTxTwoObj")
    title_standard(s)
    cw = (W - 2 * M - 1.2) / 2
    for i, (x, head, tint) in enumerate(((M, MAROON, MIST), (M + cw + 1.2, BLUE, BLUE_TINT))):
        s.ph("body", 13 + i, x, BODY_Y + 0.15, cw, 0.95, "Block title", size=16, color=WHITE,
             font=HEAD, anchor="ctr", box=head, ins=(0.35, 0.1, 0.35, 0.1),
             name="Block title")
        s.ph("obj", 1 + i, x, BODY_Y + 1.1, cw, BODY_B - BODY_Y - 1.4, "Block text",
             size=18, color=INK, box=tint, ins=(0.35, 0.3, 0.35, 0.3), name="Block body")
    footer_placeholders(s)


def layout_title_only(prs):
    s = new_layout(prs, "Title Only", "titleOnly")
    title_standard(s)
    footer_placeholders(s)


def layout_figures(prs):
    s = new_layout(prs, "Key Figures")
    title_with_subtitle(s)
    cw = (W - 2 * M) / 3
    for i in range(3):
        x = M + i * cw
        s.ph("body", 13 + i, x, BODY_Y, cw - 0.6, 2.2, "00", size=54, color=MAROON, font=HEAD,
             anchor="b", name="Figure")
        s.rect(x, BODY_Y + 2.4, 2.3, 0.15, GOLD, name="Figure rule")
        s.ph("body", 16 + i, x, BODY_Y + 2.75, cw - 0.6, 0.7, "Label", size=14, color=SLATE,
             name="Figure label")
    s.brand_rule(M, BODY_Y + 3.95, W - 2 * M, scale=0.86)
    s.ph("obj", 1, M, BODY_Y + 4.5, W - 2 * M, BODY_B - BODY_Y - 4.5, "Click to add text",
         bullets=True)
    footer_placeholders(s)


def layout_picture_caption(prs):
    s = new_layout(prs, "Picture with Caption", "picTx")
    title_standard(s)
    pw = 0.55 * (W - 2 * M)
    s.ph("pic", 13, M, BODY_Y + 0.35, pw, 10.3, "Insert picture")
    s.ph("body", 14, M, BODY_Y + 10.85, pw, 0.8, "Figure 1  |  Caption", size=12, color=SLATE,
         algn="ctr", name="Caption")
    s.ph("obj", 1, M + pw + 1.3, BODY_Y + 0.35, W - 2 * M - pw - 1.3, 10.3,
         "Click to add text", size=18, anchor="ctr")
    footer_placeholders(s)


def layout_blank(prs):
    # footer band only: no crest or title rule without a title
    s = new_layout(prs, "Blank", "blank", show_master=False, bg=WHITE)
    footer(s)


def footer_placeholders(s):
    s.ph("ftr", 11, W / 2 - 7, BAND_Y, 14, BAND, size=9, color=FOOT_TITLE, algn="ctr",
         anchor="ctr")
    s.ph("sldNum", 12, W - M - 3, BAND_Y, 3, BAND, size=9, color=WHITE, algn="r", anchor="ctr")


def ghost_crest(s, height_frac, cx=None, right=None, alpha=9):
    ch = height_frac * H
    cw = ch * LOGO_RATIO
    x = (W + right - cw) if right is not None else cx - cw / 2
    s.pic(LOGO_WHITE, x, (H - ch) / 2, cw, ch, alpha=alpha, name="Crest watermark")


def layout_section(prs):
    s = new_layout(prs, "Section Header", "secHead", show_master=False, bg=MAROON_DARK)
    ghost_crest(s, 1.22, right=2.75, alpha=9)
    s.ph("body", 13, M, 5.9, 0.7 * W, 0.7, "Section 1", size=14, color=GOLD, font=LABEL,
         spc=200, caps=True, anchor="b", name="Section number")
    s.ph("title", None, M, 7.0, 0.7 * W, 3.7, "Section title", size=36, color=WHITE,
         anchor="ctr", bullets=False)
    s.rect(M, 11.64, 5.5, 0.164, GOLD, name="Gold rule")


def layout_statement(prs):
    s = new_layout(prs, "Statement", show_master=False, bg=MAROON_DARK)
    ghost_crest(s, 1.15, cx=0.88 * W, alpha=7)
    s.ph("title", None, 0.1 * W, 2.5, 0.8 * W, H - 5, "A single statement, set large",
         size=36, color=WHITE, algn="ctr", anchor="ctr", lnspc=100)


def layout_quote(prs):
    s = new_layout(prs, "Quote", show_master=False, bg=MIST)
    s.text(M, 1.1, 6, 7, [([("“", dict(size=230, color=MAROON, font=SERIF))], "l", 0)],
           name="Quotation mark")
    # 18% maroon: set as transparency on the run's fill
    q = s.tree[-1]
    for clr in q.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"):
        clr.append(parse_xml('<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/'
                             '2006/main" val="18000"/>'))
    s.ph("body", 13, (W - 0.78 * W) / 2, 2.8, 0.78 * W, 10.4, "Quotation", size=40,
         color=MAROON_DARK, font=SERIF, italic=True, anchor="ctr", lnspc=100,
         name="Quotation")
    s.rect(M, 14.3, 3.4, 0.14, GOLD, name="Gold rule")
    s.ph("body", 14, M, 14.7, 0.8 * W, 1.17, "Attribution", size=16, color=SLATE,
         bold=True, anchor="b", name="Attribution")


def layout_photo(prs):
    s = new_layout(prs, "Photo", show_master=False, bg=MAROON_DARK)
    s.ph("pic", 13, 0, 0, W, H, "Insert a photograph")
    s.ph("body", 14, 0, H - 3.18, W, 3.18, "Caption", size=18, color=WHITE, anchor="ctr",
         box=MAROON_DARK, alpha=82, ins=(M, 0.2, M, 0.2), name="Caption")


def layout_closing(prs):
    s = new_layout(prs, "Closing", show_master=False, bg=MAROON_DARK)
    ch = 4.02
    s.pic(LOGO_WHITE, (W - ch * LOGO_RATIO) / 2, 1.8, ch * LOGO_RATIO, ch, name="Crest")
    s.ph("title", None, M, 6.2, W - 2 * M, 1.7, "Thank you", size=36, color=WHITE,
         algn="ctr", anchor="b")
    s.rect((W - 5.5) / 2, 8.3, 5.5, 0.164, GOLD, name="Gold rule")
    s.text(M, 8.85, W - 2 * M, 1.5, [
        ([(UNIVERSITY, dict(size=13, color=WHITE, font=LABEL, spc=150, caps=True))], "ctr", 0),
        ([(MOTTO, dict(size=14, color=GOLD, font=SERIF, italic=True))], "ctr", 3)],
        name="University name")
    sh = 0.85
    s.pic(SINDHI_WHITE, (W - sh * SINDHI_RATIO) / 2, 10.6, sh * SINDHI_RATIO, sh,
          name="Sindhi name")
    s.ph("body", 13, M, 12.2, W - 2 * M, 1.6, "Contact or questions line (optional)",
         size=16, color=MIST, algn="ctr", name="Contact line")
    s.text(M, H - 2.25, W - 2 * M, 0.6, [([(CONTACT, dict(size=10, color=MIST))], "ctr", 0)],
           anchor="b", name="Contact")


LAYOUTS = [layout_title_photo, layout_title_solid, layout_title_split, layout_content,
           layout_content_subtitle, layout_figures, layout_two, layout_blocks,
           layout_title_only, layout_picture_caption, layout_section, layout_statement,
           layout_quote, layout_photo, layout_closing, layout_blank]


def build_template():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    set_theme(prs)
    drop_default_layouts(prs)
    build_master(prs)
    for make in LAYOUTS:
        make(prs)
    props = prs.core_properties
    props.title = "The Shaikh Ayaz University, Shikarpur — presentation template"
    props.author = UNIVERSITY
    props.last_modified_by = UNIVERSITY   # python-pptx's stock deck names its author
    props.subject = MOTTO
    props.keywords = "SAUS; template; brand"
    return prs


# ---- Demo deck -------------------------------------------------------------------
def layout(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError(name)


def set_text(shape, paras):
    """paras: strings, lists of (text, fmt) runs, or (level, either) tuples."""
    tf = shape.text_frame
    tf.clear()
    for i, p in enumerate(paras):
        level, content = p if isinstance(p, tuple) else (0, p)
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = level
        runs = [(content, {})] if isinstance(content, str) else content
        for text, fmt in runs:
            r = para.add_run()
            r.text = text
            if fmt.get("bold"):
                r.font.bold = True
            if fmt.get("color"):
                from pptx.dml.color import RGBColor
                r.font.color.rgb = RGBColor.from_string(fmt["color"])
            if fmt.get("italic"):
                r.font.italic = True
            if fmt.get("highlight"):
                rPr = r._r.get_or_add_rPr()
                rPr.append(parse_xml(f'<a:highlight xmlns:a="http://schemas.openxmlformats.org/'
                                     f'drawingml/2006/main"><a:srgbClr val="{fmt["highlight"]}"/>'
                                     f'</a:highlight>'))
    return tf


def numbered(tf, color=MAROON):
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(emu(0.95)))
        pPr.set("indent", str(emu(-0.95)))
        for tag in ("buClr", "buSzPct", "buFont", "buChar", "buNone", "buAutoNum"):
            for e in pPr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}"):
                pPr.remove(e)
        pPr.append(parse_xml(f'<a:buClr {NS}><a:srgbClr val="{color}"/></a:buClr>'))
        pPr.append(parse_xml(f'<a:buSzPct {NS} val="100000"/>'))
        pPr.append(parse_xml(f'<a:buFont {NS} typeface="{HEAD}"/>'))
        pPr.append(parse_xml(f'<a:buAutoNum {NS} type="arabicPeriod"/>'))


def no_bullets(tf):
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        pPr.append(parse_xml(f'<a:buNone {NS}/>'))


def add_footer(slide, text):
    """Copy the layout's footer and slide-number placeholders onto the slide
    (what Insert > Header & Footer does in PowerPoint)."""
    import copy
    from pptx.enum.shapes import PP_PLACEHOLDER
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    next_id = max(int(e.get("id")) for e in slide.shapes._spTree.iter(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")) + 1
    for shape in slide.slide_layout.placeholders:
        kind = shape.placeholder_format.type
        if kind not in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            continue
        sp = copy.deepcopy(shape._element)
        sp.nvSpPr.cNvPr.set("id", str(next_id))
        next_id += 1
        sp.spPr.remove(sp.spPr.find(f"{A}xfrm"))          # position comes from the layout
        lst = sp.txBody.find(f"{A}lstStyle")
        for child in list(lst):
            lst.remove(child)
        if kind == PP_PLACEHOLDER.FOOTER:
            for p in sp.txBody.findall(f"{A}p"):
                sp.txBody.remove(p)
            sp.txBody.append(parse_xml(f'<a:p {NS}><a:r><a:rPr lang="en-GB"/>'
                                       f'<a:t>{escape(text)}</a:t></a:r></a:p>'))
        slide.shapes._spTree.append(sp)


def drop_empty(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            shape._element.getparent().remove(shape._element)


def build_demo(prs):
    FT = "Digital Transformation"
    title_parts = dict(
        title="Digital Transformation of Campus Services",
        subtitle="A roadmap for the 2026–2029 planning cycle",
        dept="Department of Computer Science", presenter="Dr. A. Presenter",
        occasion="Faculty Development Seminar  ·  21 September 2026")

    def title_slide(name):
        s = prs.slides.add_slide(layout(prs, name))
        s.shapes.title.text = title_parts["title"]
        s.placeholders[1].text = title_parts["subtitle"]
        s.placeholders[13].text = title_parts["dept"]
        s.placeholders[14].text = title_parts["presenter"]
        s.placeholders[15].text = title_parts["occasion"]
        return s

    def section(number, name):
        s = prs.slides.add_slide(layout(prs, "Section Header"))
        s.placeholders[13].text = f"Section {number}"
        s.shapes.title.text = name

    title_slide("Title Slide (photo)")

    s = prs.slides.add_slide(layout(prs, "Title and Content"))
    s.shapes.title.text = "Outline"
    numbered(set_text(s.placeholders[1], ["Context", "Programme of work", "The campus"]))
    add_footer(s, FT)

    section(1, "Context")

    s = prs.slides.add_slide(layout(prs, "Key Figures"))
    s.shapes.title.text = "The university at a glance"
    s.placeholders[10].text = "Established to revive the educational glory of Shikarpur"
    for idx, (fig, label) in zip((13, 14, 15), (("7", "Degree programmes"), ("2", "Faculties"),
                                                ("815", "Students enrolled"))):
        s.placeholders[idx].text = fig
        s.placeholders[idx + 3].text = label
    set_text(s.placeholders[1], [
        [("A ", {}), ("public sector general university", {"bold": True, "color": MAROON}),
         (" serving Shikarpur and upper Sindh.", {})],
        "Teaching, research and community engagement across the humanities, sciences "
        "and professional disciplines.",
        (1, "Second-level items use the crest’s royal blue."),
        (2, "Third-level items pick up the gold accent.")])
    add_footer(s, FT)

    s = prs.slides.add_slide(layout(prs, "Title Only"))
    s.shapes.title.text = "One university, three scripts"
    sh = Shapes(s.part, s.shapes._spTree)
    rows = [("English", None, 5.0), ("Sindhi", (SINDHI_MAROON, SINDHI_RATIO, 1.0), 7.1),
            ("Urdu", (URDU_MAROON, URDU_RATIO, 1.55), 9.1)]
    for label, img, y in rows:
        sh.text(M, y, 4, 1.0, [([(label, dict(size=14, color=SLATE))], "l", 0)], anchor="ctr")
        if img is None:
            sh.text(M + 4.5, y, 20, 1.0, [([(UNIVERSITY, dict(size=28, color=INK))], "l", 0)],
                    anchor="ctr")
        else:
            path, ratio, h = img
            sh.pic(path, M + 4.5, y + 0.5 - h / 2, h * ratio, h)
    sh.brand_rule(M, 11.4, W - 2 * M, scale=0.86)
    sh.text(M, 12.0, W - 2 * M, 2.0, [([(
        "The Sindhi and Urdu names are pictures (powerpoint/source/art), set in Lateef and "
        "Noto Nastaliq Urdu, so they display correctly on every computer.",
        dict(size=16, color=INK))], "l", 0)])
    add_footer(s, FT)

    s = prs.slides.add_slide(layout(prs, "Comparison (blocks)"))
    s.shapes.title.text = "Blocks and semantic colours"
    s.placeholders[13].text = "Standard block"
    s.placeholders[1].text = ("Maroon header on the crest’s mist tint. Use for definitions, "
                              "summaries and key takeaways.")
    s.placeholders[14].text = "Example block"
    set_text(s.placeholders[2], [
        "Royal blue header, sampled from the ring of the university crest.",
        [("Key terms", {"bold": True, "color": MAROON}), (" sit in maroon bold; ", {}),
         ("highlights", {"highlight": "EEDDB5"}), (" use a light gold wash. ", {}),
         ("Alerted text", {"color": RED}), (" is red.", {})]])
    add_footer(s, FT)

    section(2, "Programme of work")

    s = prs.slides.add_slide(layout(prs, "Statement"))
    s.shapes.title.text = ("Every service a student touches should work on a phone, "
                           "in Sindhi or English, on the first attempt.")

    s = prs.slides.add_slide(layout(prs, "Title and Content"))
    s.shapes.title.text = "Numbered workstreams"
    tf = set_text(s.placeholders[1], [
        "Student records and admissions portal",
        "Learning management system and lecture capture",
        "Campus network, Wi-Fi and data centre refresh",
        "Library discovery and digital repository"])
    numbered(tf)
    for phase, text in (("Phase 1", "Assessment and procurement."),
                        ("Phase 2", "Pilot with two faculties."),
                        ("Phase 3", "University-wide rollout.")):
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = phase + "   "
        r.font.bold = True
        from pptx.dml.color import RGBColor
        r.font.color.rgb = RGBColor.from_string(MAROON)
        p.add_run().text = text
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        pPr.append(parse_xml(f'<a:buNone {NS}/>'))
    tf.paragraphs[4]._p.get_or_add_pPr().insert(0, parse_xml(
        f'<a:spcBef {NS}><a:spcPts val="2400"/></a:spcBef>'))
    add_footer(s, FT)

    s = prs.slides.add_slide(layout(prs, "Title Only"))
    s.shapes.title.text = "Tabular data"
    data = [("Workstream", "Owner", "Year"), ("Admissions portal", "Registrar’s Office", "2026"),
            ("Learning management", "Faculty of Education", "2027"),
            ("Network refresh", "IT Directorate", "2027"),
            ("Digital repository", "Central Library", "2028")]
    tw = 18.0
    gt = s.shapes.add_table(len(data), 3, Emu(emu((W - tw) / 2)), Emu(emu(5.0)),
                            Emu(emu(tw)), Emu(emu(0.95 * len(data))))
    table = gt.table
    for i, w in enumerate((7.0, 7.8, 3.2)):
        table.columns[i].width = Emu(emu(w))
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    tblPr = table._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    for e in list(tblPr):
        tblPr.remove(e)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if c == 2 else PP_ALIGN.LEFT
            font = p.runs[0].font
            font.size = Emu(18 * 12700)
            font.bold = r == 0
            font.color.rgb = RGBColor.from_string(MAROON if r == 0 else INK)
            cell.margin_left = cell.margin_right = Emu(emu(0.15))
            tcPr = cell._tc.get_or_add_tcPr()
            top = 1.5 if r == 0 else (0.75 if r == 1 else 0)
            bottom = 1.5 if r == len(data) - 1 else 0
            borders = ""
            for tag, width in (("lnL", 0), ("lnR", 0), ("lnT", top), ("lnB", bottom)):
                if width:
                    borders += (f'<a:{tag} {NS} w="{int(width * 12700)}"><a:solidFill>'
                                f'<a:srgbClr val="{INK}"/></a:solidFill></a:{tag}>')
                else:
                    borders += f'<a:{tag} {NS} w="0"><a:noFill/></a:{tag}>'
            for b in parse_xml(f'<a:x {NS}>{borders}</a:x>'):
                tcPr.append(b)
            tcPr.append(parse_xml(f'<a:noFill {NS}/>'))
    sh = Shapes(s.part, s.shapes._spTree)
    sh.text(M, 5.0 + 0.95 * len(data) + 0.5, W - 2 * M, 0.8, [(
        [("Table 1", dict(size=14, color=MAROON, bold=True)),
         ("  |  Indicative delivery schedule.", dict(size=14, color=INK))], "ctr", 0)])
    add_footer(s, FT)

    section(3, "The campus")

    s = prs.slides.add_slide(layout(prs, "Photo"))
    s.placeholders[13].insert_picture(os.path.join(ASSETS, "saus-campus-gate.jpg"))
    s.placeholders[14].text = ("The campus on Main Road, Shikarpur — the setting for every "
                               "service described in this plan.")

    s = prs.slides.add_slide(layout(prs, "Quote"))
    s.placeholders[13].text = "Revival of Educational Glory of Shikarpur."
    s.placeholders[14].text = "Motto of The Shaikh Ayaz University, Shikarpur"

    s = prs.slides.add_slide(layout(prs, "Picture with Caption"))
    s.shapes.title.text = "Figures"
    s.placeholders[13].insert_picture(os.path.join(ASSETS, "saus-monument.jpg"))
    set_text(s.placeholders[14], [[("Figure 1", {"bold": True, "color": MAROON}),
                                   ("  |  The campus monument.", {})]])
    set_text(s.placeholders[1], [
        "Captions are numbered and the label is set in maroon.",
        "Photographs sit flush to the text column; full-bleed images get their own slide "
        "with the Photo layout."])
    for p in s.placeholders[1].text_frame.paragraphs:
        p.space_before = Emu(12 * 12700)
    add_footer(s, FT)

    s = prs.slides.add_slide(layout(prs, "Closing"))
    s.shapes.title.text = "Thank you"
    drop_empty(s, 13)

    # the other two title designs
    title_slide("Title Slide (solid)")
    title_slide("Title Slide (split)")

    props = prs.core_properties
    props.title = title_parts["title"]


def save_potx(prs, path):
    """Save as .pptx, then relabel the main part so PowerPoint opens it as a template."""
    tmp = path + ".tmp"
    prs.save(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml")
            dst.writestr(item, data)
    os.remove(tmp)


if __name__ == "__main__":
    save_potx(build_template(), os.path.join(OUT, "SAUS-template.potx"))
    demo = build_template()
    build_demo(demo)
    demo.save(os.path.join(OUT, "SAUS-slides-demo.pptx"))
    print("wrote SAUS-template.potx and SAUS-slides-demo.pptx")
