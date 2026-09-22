"""Build the SAUS Word letterhead templates and demo letter.

    python build_docx.py

Writes, in the folder above this one:
  SAUS-letterhead.dotx      any office: department and e-mail are fill-in fields
  SAUS-letterhead-CS.dotx   Department of Computer Science, hod.cs@saus.edu.pk
  SAUS-letter-demo.docx     the LaTeX demo letter (latex/saus-letter-demo.tex)

The layout follows latex/saus-letter.cls: A4, 22 mm side margins, the
letterhead on page 1, "Ref. No. ... Dated ..." on one line above the
addressee, the contact footer on page 1 and a slim running head on later pages.

Needs python-docx and Pillow (pip install python-docx pillow).
"""
import copy
import os
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
ASSETS = os.path.join(ROOT, "latex", "assets")
ART = os.path.join(ROOT, "powerpoint", "source", "art")
OUT = os.path.normpath(os.path.join(HERE, ".."))

MAROON, MAROON_LIGHT, GOLD = "6B2A4C", "8E4166", "C39B3E"
INK, SLATE, RULE = "231C21", "6A626B", "D9C9D2"
SANS, SEMIBOLD, MEDIUM, SERIF, MONO = ("Fira Sans", "Fira Sans SemiBold", "Fira Sans Medium",
                                       "EB Garamond", "Fira Mono")
UNIVERSITY = "The Shaikh Ayaz University, Shikarpur"
MOTTO = "Revival of Educational Glory of Shikarpur"
ADDRESS, PHONE, WEB = "Main Road, Shikarpur, Sindh", "+92 726 512039", "saus.edu.pk"
SEP = " · "
TEXT_W = 166                              # mm between the 22 mm margins
W = nsdecls("w")


# ---- Low-level helpers ---------------------------------------------------------
def rpr_xml(size=None, color=None, font=None, bold=False, italic=False, caps=False,
            spacing=None, style=None):
    x = f'<w:rStyle w:val="{style}"/>' if style else ""
    if font:
        x += f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>'
    x += "<w:b/>" if bold else ""
    x += "<w:i/>" if italic else ""
    x += "<w:caps/>" if caps else ""
    x += f'<w:color w:val="{color}"/>' if color else ""
    x += f'<w:spacing w:val="{spacing}"/>' if spacing else ""
    x += f'<w:sz w:val="{int(size * 2)}"/><w:szCs w:val="{int(size * 2)}"/>' if size else ""
    return f"<w:rPr>{x}</w:rPr>"


def run(text, **fmt):
    t = escape(text)
    parts = []
    for i, chunk in enumerate(t.split("\t")):
        if i:
            parts.append("<w:tab/>")
        if chunk:
            parts.append(f'<w:t xml:space="preserve">{chunk}</w:t>')
    return parse_xml(f"<w:r {W}>{rpr_xml(**fmt)}{''.join(parts)}</w:r>")


def add_runs(p, *runs):
    for r in runs:
        p._p.append(run(r[0], **r[1]) if isinstance(r, tuple) else r)


def field(instr, cached, **fmt):
    """A complex field as a list of runs."""
    rp = rpr_xml(**fmt)
    return [parse_xml(f'<w:r {W}>{rp}<w:fldChar w:fldCharType="begin"/></w:r>'),
            parse_xml(f'<w:r {W}>{rp}<w:instrText xml:space="preserve"> {escape(instr)} </w:instrText></w:r>'),
            parse_xml(f'<w:r {W}>{rp}<w:fldChar w:fldCharType="separate"/></w:r>'),
            parse_xml(f'<w:r {W}>{rp}<w:t xml:space="preserve">{escape(cached)}</w:t></w:r>'),
            parse_xml(f'<w:r {W}>{rp}<w:fldChar w:fldCharType="end"/></w:r>')]


def ref_field(cached, sep, **fmt):
    """{ IF "{ STYLEREF "SAUS Ref Value" }" = "Error*" "" "{ STYLEREF ... } · " }:
    the reference number followed by a separator, or nothing when the letter
    has no reference number."""
    rp = rpr_xml(**fmt)

    def instr(t):
        return parse_xml(f'<w:r {W}>{rp}<w:instrText xml:space="preserve">{escape(t)}</w:instrText></w:r>')

    def char(kind):
        return parse_xml(f'<w:r {W}>{rp}<w:fldChar w:fldCharType="{kind}"/></w:r>')

    def styleref():
        return field('STYLEREF "SAUS Ref Value"', cached, **fmt)

    result = cached + sep if cached else ""
    # inner: { IF "{ref}" = "Error*" "" "{ref} · " }   (no text in that style)
    inner = [char("begin"), instr(' IF "')] + styleref() + [instr('" = "Error*" "" "')]
    inner += styleref() + [instr(f'{sep}" '), char("separate")]
    inner += [parse_xml(f'<w:r {W}>{rp}<w:t xml:space="preserve">{escape(result)}</w:t></w:r>'),
              char("end")]
    # outer: { IF "{ref}" = "?*" "{inner}" "" }        (empty reference)
    out = [char("begin"), instr(' IF "')] + styleref() + [instr('" = "?*" "')]
    out += inner + [instr('" "" '), char("separate")]
    out += [parse_xml(f'<w:r {W}>{rp}<w:t xml:space="preserve">{escape(result)}</w:t></w:r>'),
            char("end")]
    return out


def hyperlink(p, url, text, **fmt):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = parse_xml(f'<w:hyperlink {W} {nsdecls("r")} r:id="{rid}"/>')
    h.append(run(text, **fmt))
    p._p.append(h)


def sdt_inline(p, alias, tag, text, placeholder=False, date=None, style=None, **fmt):
    """Inline content control (plain text, or a date picker when date is given)."""
    kind = "<w:text/>"
    if date is not None:
        full = f' w:fullDate="{date}T00:00:00Z"' if date else ""
        kind = (f'<w:date{full}><w:dateFormat w:val="d MMMM yyyy"/><w:lid w:val="en-GB"/>'
                f'<w:storeMappedDataAs w:val="dateTime"/><w:calendar w:val="gregorian"/></w:date>')
    sp = f'<w:rPr><w:rStyle w:val="{style}"/></w:rPr>' if style else ""
    ph = "<w:showingPlcHdr/>" if placeholder else ""
    r = run(text, style="PlaceholderText" if placeholder else style, **fmt)
    sdt = parse_xml(f'<w:sdt {W}><w:sdtPr>{sp}<w:alias w:val="{alias}"/><w:tag w:val="{tag}"/>'
                    f'{ph}{kind}</w:sdtPr><w:sdtContent/></w:sdt>')
    sdt.find(qn("w:sdtContent")).append(r)
    p._p.append(sdt)


def sdt_block(paragraphs, alias, tag, placeholder=False):
    """Wrap consecutive body paragraphs in a rich-text content control."""
    first = paragraphs[0]._p
    ph = "<w:showingPlcHdr/>" if placeholder else ""
    sdt = parse_xml(f'<w:sdt {W}><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{tag}"/>'
                    f'{ph}</w:sdtPr><w:sdtContent/></w:sdt>')
    first.addprevious(sdt)
    content = sdt.find(qn("w:sdtContent"))
    for p in paragraphs:
        content.append(p._p)
        if placeholder:
            for r in p._p.iter(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = parse_xml(f"<w:rPr {W}/>")
                    r.insert(0, rpr)
                rs = rpr.find(qn("w:rStyle"))
                if rs is None:
                    rpr.insert(0, parse_xml(f'<w:rStyle {W} w:val="PlaceholderText"/>'))


def border(p, side, color=RULE, size=4, space=4):
    ppr = p._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = parse_xml(f"<w:pBdr {W}/>")
        ppr.append(bdr)
    bdr.append(parse_xml(f'<w:{side} {W} w:val="single" w:sz="{size}" w:space="{space}" '
                         f'w:color="{color}"/>'))


def brand_rule_png():
    """The maroon-gold-hairline signature rule at 11 pt, as a 600 dpi image."""
    path = os.path.join(HERE, "brand-rule.png")
    dpmm = 600 / 25.4
    w, h = round(TEXT_W * dpmm), round(0.78 * dpmm)
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    m, g = round(5.67 * dpmm), round(2.48 * dpmm)
    d.rectangle([0, 0, m - 1, h - 1], fill="#" + MAROON)
    d.rectangle([m, 0, m + g - 1, h - 1], fill="#" + GOLD)
    hl = max(2, round(0.21 * dpmm))
    top = (h - hl) // 2
    d.rectangle([m + g, top, w - 1, top + hl - 1], fill="#" + RULE)
    img.save(path, dpi=(600, 600))
    return path


# ---- Styles ----------------------------------------------------------------------
def setup_styles(doc):
    styles = doc.styles
    # document defaults: Fira Sans 11 pt, British English, no theme fonts
    rpr = styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for child in list(rpr):
        rpr.remove(child)
    rpr.append(parse_xml(f'<w:rFonts {W} w:ascii="{SANS}" w:hAnsi="{SANS}" w:eastAsia="{SANS}" '
                         f'w:cs="Lateef"/>'))
    rpr.append(parse_xml(f'<w:color {W} w:val="{INK}"/>'))
    rpr.append(parse_xml(f'<w:sz {W} w:val="22"/>'))
    rpr.append(parse_xml(f'<w:szCs {W} w:val="22"/>'))
    rpr.append(parse_xml(f'<w:lang {W} w:val="en-GB" w:eastAsia="en-GB" w:bidi="ur-PK"/>'))

    normal = styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(6.5)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    def para(name, base="Normal", **kw):
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base]
        st.quick_style = True
        f, p = st.font, st.paragraph_format
        for k, v in kw.items():
            if k in ("size",):
                f.size = Pt(v)
            elif k == "color":
                f.color.rgb = RGBColor.from_string(v)
            elif k == "font":
                f.name = v
            elif k in ("bold", "italic", "all_caps"):
                setattr(f, k, v)
            elif k in ("before", "after"):
                setattr(p, "space_" + k, Pt(v))
            else:
                setattr(p, k, v)
        return st

    def char(name, **kw):
        st = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        st.quick_style = True
        for k, v in kw.items():
            if k == "size":
                st.font.size = Pt(v)
            elif k == "color":
                st.font.color.rgb = RGBColor.from_string(v)
            elif k == "font":
                st.font.name = v
            else:
                setattr(st.font, k, v)
        return st

    hdr = para("SAUS Header", before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    hdr.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT)   # right
    ref = para("SAUS Reference", before=0, after=31, alignment=WD_ALIGN_PARAGRAPH.LEFT,
               keep_with_next=True)
    ref.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT)
    para("SAUS Address", before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT,
         keep_with_next=True, keep_together=True)
    para("SAUS Subject", font=SEMIBOLD, color=MAROON, before=34, after=24,
         alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    para("SAUS Opening", before=0, after=13, alignment=WD_ALIGN_PARAGRAPH.LEFT,
         keep_with_next=True)
    para("SAUS Table Text", before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    para("SAUS Closing", before=12, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT,
         keep_with_next=True)
    para("SAUS Signature", before=45, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT,
         keep_together=True)
    para("SAUS Enclosure", before=6, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    para("SAUS Notice Title", font=SEMIBOLD, color=MAROON, size=15, before=6, after=18,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)

    char("SAUS Label", font=MEDIUM, color=MAROON, size=10)
    char("SAUS Ref Value")
    char("SAUS Key Term", bold=True, color=MAROON)
    char("SAUS Code", font=MONO, size=10)
    char("Placeholder Text", color="808080")

    # Lists: maroon squares, then blue and gold; numbers in maroon
    numbering = doc.part.numbering_part.element
    bullet = (f'<w:abstractNum {W} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>')
    for lvl, color in enumerate((MAROON, "1A1A91", GOLD)):
        bullet += (f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
                   f'<w:lvlText w:val=""/><w:lvlJc w:val="left"/><w:pPr>'
                   f'<w:ind w:left="{340 * (lvl + 1)}" w:hanging="340"/></w:pPr><w:rPr>'
                   f'<w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/>'
                   f'<w:color w:val="{color}"/><w:sz w:val="{18 - 2 * lvl}"/></w:rPr></w:lvl>')
    bullet += "</w:abstractNum>"
    number = (f'<w:abstractNum {W} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
              f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
              f'<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="397" '
              f'w:hanging="397"/></w:pPr><w:rPr><w:rFonts w:ascii="{SEMIBOLD}" w:hAnsi="{SEMIBOLD}"/>'
              f'<w:color w:val="{MAROON}"/></w:rPr></w:lvl></w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    for xml in (bullet, number):
        el = parse_xml(xml)
        if first_num is not None:
            first_num.addprevious(el)
        else:
            numbering.append(el)
    numbering.append(parse_xml(f'<w:num {W} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    numbering.append(parse_xml(f'<w:num {W} w:numId="91"><w:abstractNumId w:val="91"/></w:num>'))
    for name, num_id in (("List Bullet", 90), ("List Number", 91)):
        st = styles[name]
        ppr = st.element.get_or_add_pPr()
        for old in ppr.findall(qn("w:numPr")):
            ppr.remove(old)
        ppr.insert(0, parse_xml(f'<w:numPr {W}><w:numId w:val="{num_id}"/></w:numPr>'))
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.quick_style = True

    # Booktabs-style table: rules above and below, a thinner one under the header
    styles.element.append(parse_xml(
        f'<w:style {W} w:type="table" w:customStyle="1" w:styleId="SAUSTable">'
        f'<w:name w:val="SAUS Table"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="59"/>'
        f'<w:qFormat/><w:pPr><w:spacing w:before="0" w:after="0"/><w:jc w:val="left"/></w:pPr>'
        f'<w:tblPr><w:jc w:val="center"/><w:tblBorders>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="{INK}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{INK}"/></w:tblBorders>'
        f'<w:tblCellMar><w:top w:w="45" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        f'<w:bottom w:w="45" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar></w:tblPr>'
        f'<w:tblStylePr w:type="firstRow"><w:rPr><w:rFonts w:ascii="{SEMIBOLD}" w:hAnsi="{SEMIBOLD}"/>'
        f'</w:rPr><w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="4" w:space="0" '
        f'w:color="{INK}"/></w:tcBorders></w:tcPr></w:tblStylePr></w:style>'))


def setup_settings(doc):
    settings = doc.settings.element
    tab = settings.find(qn("w:defaultTabStop"))
    hyph = [parse_xml(f"<w:autoHyphenation {W}/>"),
            parse_xml(f'<w:consecutiveHyphenLimit {W} w:val="2"/>')]
    anchor = tab
    for el in hyph:
        if anchor is not None:
            anchor.addnext(el)
            anchor = el
        else:
            settings.insert(0, el)


# ---- Page furniture -----------------------------------------------------------------
def letterhead(section, department, rule_png):
    """Page-1 header: crest, name, motto, office; Sindhi and Urdu names; the rule."""
    header = section.first_page_header
    header.is_linked_to_previous = False
    table = header.add_table(1, 3, Mm(TEXT_W))
    header._element.remove(header.paragraphs[0]._p)          # the empty default paragraph
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(parse_xml(f'<w:tblLayout {W} w:type="fixed"/>'))
    tbl_pr.append(parse_xml(f'<w:tblCellMar {W}><w:left w:w="0" w:type="dxa"/>'
                            f'<w:right w:w="0" w:type="dxa"/></w:tblCellMar>'))
    widths = (19, TEXT_W - 19 - 49, 49)
    for cell, w in zip(table.rows[0].cells, widths):
        cell.width = Mm(w)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 56.6929)))

    crest, text, native = table.rows[0].cells
    p = crest.paragraphs[0]
    p.style = "SAUS Header"
    p.add_run().add_picture(os.path.join(ASSETS, "saus-logo.png"), width=Mm(19))

    p = text.paragraphs[0]
    p.style = "SAUS Header"
    p.paragraph_format.left_indent = Mm(5)
    add_runs(p, (UNIVERSITY, dict(size=11.5, color=MAROON, font=SEMIBOLD, caps=True,
                                  spacing=20)))
    p = text.add_paragraph(style="SAUS Header")
    p.paragraph_format.left_indent = Mm(5)
    p.paragraph_format.space_before = Pt(2.5)
    add_runs(p, (MOTTO, dict(size=11.5, color=MAROON_LIGHT, font=SERIF, italic=True)))
    p = text.add_paragraph(style="SAUS Header")
    p.paragraph_format.left_indent = Mm(5)
    p.paragraph_format.space_before = Pt(5)
    fmt = dict(size=8.5, color=SLATE, font=MEDIUM, caps=True, spacing=20)
    if department:
        add_runs(p, (department, fmt))
    else:
        sdt_inline(p, "Department or office", "department", "Department or office",
                   placeholder=True, **fmt)

    p = native.paragraphs[0]
    p.style = "SAUS Header"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(os.path.join(ART, "sindhi-maroon.png"), width=Mm(41.4))
    p = native.add_paragraph(style="SAUS Header")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    p.add_run().add_picture(os.path.join(ART, "urdu-maroon.png"), width=Mm(36.3))

    p = header.add_paragraph(style="SAUS Header")
    p.paragraph_format.space_before = Mm(3.5)
    p.add_run().add_picture(rule_png, width=Mm(TEXT_W))


def contact_footer(section, email):
    footer = section.first_page_footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.style = "SAUS Header"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    border(p, "top", space=6)
    fmt = dict(size=8.5, color=SLATE)
    add_runs(p, (ADDRESS + SEP + PHONE + SEP, fmt))
    hyperlink(p, "https://" + WEB, WEB, size=8.5, color=MAROON)
    add_runs(p, (SEP, fmt))
    if email:
        hyperlink(p, "mailto:" + email, email, **fmt)
    else:
        sdt_inline(p, "Office e-mail", "email", "office@saus.edu.pk", placeholder=True, **fmt)


def running_head(section, cached_ref):
    """Later pages: small crest and name; reference number and page count."""
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.style = "SAUS Header"
    r = p.add_run()
    r.add_picture(os.path.join(ASSETS, "saus-logo.png"), height=Mm(8))
    r._r.get_or_add_rPr().append(parse_xml(f'<w:position {W} w:val="-8"/>'))
    add_runs(p, (" " + UNIVERSITY, dict(size=9, color=MAROON, font=MEDIUM, caps=True,
                                             spacing=20)), ("\t", {}))
    fmt = dict(size=9, color=SLATE)
    for r in ref_field(cached_ref, SEP, **fmt):
        p._p.append(r)
    add_runs(p, ("Page ", fmt))
    for r in field("PAGE", "2", **fmt):
        p._p.append(r)
    add_runs(p, (" of ", fmt))
    for r in field("NUMPAGES", "2", **fmt):
        p._p.append(r)
    border(p, "bottom", size=4, space=4)
    section.footer.is_linked_to_previous = False


# ---- Letter body ------------------------------------------------------------------------
def reference_line(doc, ref, date, ref_placeholder):
    p = doc.add_paragraph(style="SAUS Reference")
    add_runs(p, ("Ref. No.", dict(style="SAUSLabel")), (" ", {}))
    if ref:
        add_runs(p, (ref, dict(style="SAUSRefValue")))
    else:
        sdt_inline(p, "Reference number", "ref", ref_placeholder, placeholder=True,
                   style="SAUSRefValue")
    add_runs(p, ("\t", {}), ("Dated", dict(style="SAUSLabel")), (" ", {}))
    if date:
        sdt_inline(p, "Date", "date", date[1], date=date[0])
    else:
        sdt_inline(p, "Date", "date", "Select date", placeholder=True, date="")


def paragraphs(doc, lines, style, placeholder=False, first_bold=False):
    out = []
    for i, line in enumerate(lines):
        p = doc.add_paragraph(style=style)
        add_runs(p, (line, dict(bold=first_bold and i == 0)))
        out.append(p)
    return out


def build(kind):
    """kind: 'template' (any office), 'cs' (Computer Science template) or 'demo'."""
    doc = Document()
    setup_styles(doc)
    setup_settings(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin, sec.bottom_margin = Mm(34), Mm(28)
    sec.header_distance, sec.footer_distance = Mm(12), Mm(12)
    sec.different_first_page_header_footer = True

    cs = kind in ("cs", "demo")
    department = "Department of Computer Science" if cs else None
    email = "hod.cs@saus.edu.pk" if cs else None
    demo = kind == "demo"
    rule = brand_rule_png()
    letterhead(sec, department, rule)
    contact_footer(sec, email)
    running_head(sec, "SAUS/CS/2026/01" if demo else "")

    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):              # python-docx's empty first paragraph
        body.remove(p)

    reference_line(doc, "SAUS/CS/2026/01" if demo else None,
                   ("2026-09-21", "21 September 2026") if demo else None,
                   "SAUS/CS/2026/..." if cs else "SAUS/.../2026/...")

    if demo:
        paragraphs(doc, ["All Faculty Members", "Department of Computer Science", UNIVERSITY],
                   "SAUS Address", first_bold=True)
    else:
        ps = paragraphs(doc, ["Recipient name", "Designation, organisation", "Address"],
                        "SAUS Address", first_bold=True)
        sdt_block(ps, "Addressee", "addressee", placeholder=True)

    p = doc.add_paragraph(style="SAUS Subject")
    add_runs(p, ("Subject: ", {}))
    if demo:
        add_runs(p, ("University templates for lectures, posters and letters", {}))
    else:
        sdt_inline(p, "Subject", "subject", "Subject of the letter", placeholder=True)

    if demo:
        demo_body(doc)
    else:
        p = doc.add_paragraph(style="SAUS Opening")
        add_runs(p, ("Dear Sir/Madam,", {}))
        ps = paragraphs(doc, ["Type the letter here. Use the List Bullet and List Number "
                              "styles for lists and SAUS Table for tables."], "Normal")
        sdt_block(ps, "Letter", "body", placeholder=True)

    p = doc.add_paragraph(style="SAUS Closing")
    add_runs(p, ("Yours sincerely,", {}))
    p = doc.add_paragraph(style="SAUS Signature")
    if cs:
        add_runs(p, ("Head of Department", {}))
        p.add_run().add_break(WD_BREAK.LINE)
        add_runs(p, ("Department of Computer Science", {}))
    else:
        sdt_inline(p, "Signatory", "signatory", "Name, designation", placeholder=True)

    enc = doc.add_paragraph(style="SAUS Enclosure")
    enc.paragraph_format.space_before = Pt(22)
    add_runs(enc, ("Enclosure:\t", {}))
    cc = doc.add_paragraph(style="SAUS Enclosure")
    # labels of different widths: continuation lines align under the first entry
    for par, width in ((enc, 21), (cc, 16)):
        pf = par.paragraph_format
        pf.left_indent, pf.first_line_indent = Mm(width), Mm(-width)
        pf.tab_stops.add_tab_stop(Mm(width))
    add_runs(cc, ("Copy to:\t", {}))
    if demo:
        add_runs(enc, ("Template guide (README)", {}))
        add_runs(cc, ("IT Directorate", {}))
        cc.add_run().add_break(WD_BREAK.LINE)
        add_runs(cc, ("Office file", {}))
    else:
        sdt_inline(enc, "Enclosure", "enclosure", "Enclosures, or delete this line",
                   placeholder=True)
        sdt_inline(cc, "Copy to", "cc", "Copies, or delete this line", placeholder=True)

    props = doc.core_properties
    props.author = UNIVERSITY
    props.last_modified_by = UNIVERSITY   # not whoever last opened it in Word
    props.title = ("Letter: University templates" if demo else
                   "The Shaikh Ayaz University, Shikarpur — letterhead")
    props.subject = MOTTO
    props.comments = "Generated by word/source/build_docx.py"
    return doc


def demo_body(doc):
    """The circular from latex/saus-letter-demo.tex, adapted to mention the Office files."""
    p = doc.add_paragraph(style="SAUS Opening")
    add_runs(p, ("Dear Colleagues,", {}))
    doc.add_paragraph(
        "The university’s IT Directorate has prepared a set of document templates that "
        "give our presentations, research posters and correspondence one consistent "
        "identity. They are built on the university crest, use the same colours and "
        "typefaces throughout, and set the university’s name in English, Sindhi and "
        "Urdu. I would like the department to adopt them for its teaching, research and "
        "official work.")
    doc.add_paragraph("The templates cover four kinds of document:")
    for lead, rest in (("Slides", " — a presentation theme for lectures and seminars, in "
                                  "LaTeX and PowerPoint."),
                       ("Posters", " — research posters from A3 to A0, in portrait or "
                                   "landscape, for conferences and the department’s "
                                   "research events."),
                       ("Letters", " — this letterhead, in LaTeX and Word, for the "
                                   "department’s official correspondence."),
                       ("Certificates", " — certificates of participation and "
                                        "achievement, made in batches from a list of names.")):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, (lead, dict(bold=True)), (rest, {}))
    doc.add_paragraph("All of them share one set of colours, typefaces and details, so every "
                      "document the department produces looks the same.")

    rows = [("Template", "LaTeX", "Office"),
            ("Slides", "saus-slides-starter.tex", "SAUS-template.potx"),
            ("Posters", "saus-poster-starter.tex", "—"),
            ("Letters", "saus-letter-starter.tex", "SAUS-letterhead.dotx"),
            ("Certificates", "saus-certificate-starter.tex", "—")]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = doc.styles["SAUS Table"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (28, 66, 47)
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 56.6929)))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Mm(widths[c])
            p = cell.paragraphs[0]
            p.style = "SAUS Table Text"
            code = r > 0 and c > 0 and value != "—"
            fmt = dict(style="SAUSCode") if code else (dict(font=SEMIBOLD) if r == 0 else {})
            add_runs(p, (value, fmt))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(8)

    doc.add_paragraph(
        "The templates run on any standard LaTeX installation, including Overleaf, and "
        "in Microsoft PowerPoint and Word once the university fonts are installed. "
        "Colleagues who have not used them before can start from the starter files, which "
        "contain every field they need to fill in.")
    doc.add_paragraph(
        "Please use the slide template for lectures and seminars from the coming semester, "
        "and the poster template for any research you present on behalf of the department. "
        "Course coordinators are requested to share this letter with the visiting faculty "
        "and teaching assistants in their courses.")
    doc.add_paragraph(
        "The IT Directorate is glad to help anyone setting up the templates. Questions about "
        "their use in the department may be sent to me at the address below.")


def save(doc, path, template):
    doc.save(path)
    if not template:
        return
    tmp = path + ".tmp"
    os.replace(path, tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml")
            dst.writestr(item, data)
    os.remove(tmp)


if __name__ == "__main__":
    save(build("template"), os.path.join(OUT, "SAUS-letterhead.dotx"), True)
    save(build("cs"), os.path.join(OUT, "SAUS-letterhead-CS.dotx"), True)
    save(build("demo"), os.path.join(OUT, "SAUS-letter-demo.docx"), False)
    os.remove(os.path.join(HERE, "brand-rule.png"))
    print("wrote SAUS-letterhead.dotx, SAUS-letterhead-CS.dotx and SAUS-letter-demo.docx")
