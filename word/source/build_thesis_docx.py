"""Build the SAUS Word thesis templates and demo report.

    python build_thesis_docx.py
    powershell -File finalize.ps1 ../SAUS-thesis-demo.docx   (updates contents and fields)

Writes, in the folder above this one:
  SAUS-thesis-FYP.dotx    final year project report (BS), "We" declaration
  SAUS-thesis.dotx        MS / MPhil / PhD thesis, "I" declaration
  SAUS-thesis-demo.docx   the LaTeX demo report (latex/saus-thesis-demo.tex)

The layout follows latex/saus-thesis.cls: A4, 38 mm binding margin, 1.5 line
spacing, EB Garamond text with Fira Sans headings; title page, certificate of
approval, declaration, abstracts, contents and lists in a roman-numbered front
matter; chapters numbered "Chapter 1", sections "1.1"; captions "Figure 3.1";
a running head with the chapter and the page number.

Needs python-docx and Pillow (pip install python-docx pillow).
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw

import build_docx as L
from build_docx import W, add_runs, field, run, sdt_block, sdt_inline, border

HERE, ASSETS, ART_PPT, OUT = L.HERE, L.ASSETS, L.ART, L.OUT
ART = os.path.join(HERE, "art")
MAROON, MAROON_LIGHT, MAROON_DARK, GOLD = L.MAROON, L.MAROON_LIGHT, "4A1C34", L.GOLD
INK, SLATE, RULE, MIST, BLUE = L.INK, L.SLATE, L.RULE, "F4EDF1", "1A1A91"
SANS, SEMIBOLD, MEDIUM, SERIF, MONO = L.SANS, L.SEMIBOLD, L.MEDIUM, L.SERIF, L.MONO
LIGHT = "Fira Sans Light"
UNIVERSITY, MOTTO = L.UNIVERSITY, L.MOTTO
TEXT_W = 210 - 38 - 25                      # mm
TW = int(TEXT_W * 56.6929)                  # text width in twips
SEP = " · "
LABEL_TAB = int(38 * 56.6929)               # where chapter titles start, after "CHAPTER 1"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
BUILTIN = {"toc 1": "TOC1", "toc 2": "TOC2", "toc 3": "TOC3", "toc 4": "TOC4",
           "table of figures": "TableofFigures"}


# ---- Styles ------------------------------------------------------------------------------
def style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal", font=None, size=None,
          color=None, bold=None, italic=None, caps=None, align=None, before=None, after=None,
          line=None, keep_next=None, keep_lines=None, page_break=None, left=None, first=None,
          outline=None, spacing=None, next_style=None):
    styles = doc.styles
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, kind)
    if name in BUILTIN:
        # Word only uses its own style for a TOC if name and id match its built-in ones
        st.element.name_val, st.element.styleId = name, BUILTIN[name]
        st.element.attrib.pop(qn("w:customStyle"), None)
    if kind == WD_STYLE_TYPE.PARAGRAPH and base:
        st.base_style = styles[base]
    st.quick_style = True
    f = st.font
    if font:
        f.name = font
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:cs"), font)
        # built-in styles (headings, caption) link to theme fonts, which win over names
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(qn("w:" + attr), None)
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = RGBColor.from_string(color)
    for attr, val in (("bold", bold), ("italic", italic), ("all_caps", caps)):
        if val is not None:
            setattr(f, attr, val)
    if spacing:
        st.element.get_or_add_rPr().append(parse_xml(f'<w:spacing {W} w:val="{spacing}"/>'))
    if kind != WD_STYLE_TYPE.PARAGRAPH:
        return st
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    for attr, val in (("keep_with_next", keep_next), ("keep_together", keep_lines),
                      ("page_break_before", page_break)):
        if val is not None:
            setattr(pf, attr, val)
    if left is not None:
        pf.left_indent = Mm(left)
    if first is not None:
        pf.first_line_indent = Mm(first)
    if outline is not None:
        st.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {W} w:val="{outline}"/>'))
    if next_style:
        st.next_paragraph_style = styles[next_style]
    return st


def numbering(doc):
    """Chapter, appendix, reference and list numbering."""
    num = doc.part.numbering_part.element
    wrap_tab = f'<w:tabs><w:tab w:val="left" w:pos="{LABEL_TAB}"/></w:tabs>'
    label = (f'<w:rPr><w:rFonts w:ascii="{MEDIUM}" w:hAnsi="{MEDIUM}"/><w:b w:val="0"/><w:caps/>'
             f'<w:color w:val="{GOLD}"/><w:spacing w:val="40"/><w:position w:val="10"/>'
             f'<w:sz w:val="22"/></w:rPr>')
    # "CHAPTER 1" in small gold capitals, raised beside the title; the title
    # starts at a fixed tab so every chapter lines up
    headings = (f'<w:abstractNum {W} w:abstractNumId="80"><w:multiLevelType w:val="multilevel"/>'
                f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                f'<w:pStyle w:val="Heading1"/><w:lvlText w:val="Chapter %1"/><w:lvlJc w:val="left"/>'
                f'<w:suff w:val="tab"/><w:pPr>{wrap_tab}<w:ind w:left="0" w:firstLine="0"/></w:pPr>'
                f'{label}</w:lvl>'
                f'<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                f'<w:pStyle w:val="Heading2"/><w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="737" w:hanging="737"/></w:pPr></w:lvl>'
                f'<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                f'<w:pStyle w:val="Heading3"/><w:lvlText w:val="%1.%2.%3"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="964" w:hanging="964"/></w:pPr></w:lvl></w:abstractNum>')
    appendix = (f'<w:abstractNum {W} w:abstractNumId="81"><w:multiLevelType w:val="singleLevel"/>'
                f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperLetter"/>'
                f'<w:pStyle w:val="SAUSAppendixHeading"/><w:lvlText w:val="Appendix %1"/>'
                f'<w:lvlJc w:val="left"/><w:suff w:val="tab"/><w:pPr>{wrap_tab}'
                f'<w:ind w:left="0" w:firstLine="0"/></w:pPr>{label}</w:lvl></w:abstractNum>')
    refs = (f'<w:abstractNum {W} w:abstractNumId="82"><w:multiLevelType w:val="singleLevel"/>'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            f'<w:pStyle w:val="SAUSReference"/><w:lvlText w:val="[%1]"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="624" w:hanging="624"/></w:pPr></w:lvl></w:abstractNum>')
    bullet = f'<w:abstractNum {W} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>'
    for lvl, color in enumerate((MAROON, BLUE, GOLD)):
        bullet += (f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
                   f'<w:lvlText w:val=""/><w:lvlJc w:val="left"/><w:pPr>'
                   f'<w:ind w:left="{397 * (lvl + 1)}" w:hanging="397"/></w:pPr><w:rPr>'
                   f'<w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/>'
                   f'<w:color w:val="{color}"/><w:sz w:val="{18 - 2 * lvl}"/></w:rPr></w:lvl>')
    bullet += "</w:abstractNum>"
    number = (f'<w:abstractNum {W} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
              f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
              f'<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="454" '
              f'w:hanging="454"/></w:pPr><w:rPr><w:rFonts w:ascii="{SEMIBOLD}" w:hAnsi="{SEMIBOLD}"/>'
              f'<w:color w:val="{MAROON}"/></w:rPr></w:lvl></w:abstractNum>')
    first_num = num.find(qn("w:num"))
    for xml in (headings, appendix, refs, bullet, number):
        el = parse_xml(xml)
        (first_num.addprevious(el) if first_num is not None else num.append(el))
    for i in (80, 81, 82, 90, 91):
        num.append(parse_xml(f'<w:num {W} w:numId="{i}"><w:abstractNumId w:val="{i}"/></w:num>'))


def link_numbering(doc, name, num_id, ilvl=0):
    ppr = doc.styles[name].element.get_or_add_pPr()
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    ppr.insert(0, parse_xml(f'<w:numPr {W}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'))


def setup_styles(doc):
    styles = doc.styles
    rpr = styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for child in list(rpr):
        rpr.remove(child)
    for xml in (f'<w:rFonts {W} w:ascii="{SERIF}" w:hAnsi="{SERIF}" w:eastAsia="{SERIF}" w:cs="Lateef"/>',
                f'<w:color {W} w:val="{INK}"/>', f'<w:sz {W} w:val="24"/>', f'<w:szCs {W} w:val="24"/>',
                f'<w:lang {W} w:val="en-GB" w:eastAsia="en-GB" w:bidi="ur-PK"/>'):
        rpr.append(parse_xml(xml))
    normal = styles["Normal"]
    normal.font.name, normal.font.size = SERIF, Pt(12)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(8), 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    J, C, Lf = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    head = dict(font=SEMIBOLD, color=MAROON, align=Lf, line=1.0, keep_next=True, keep_lines=True)
    h1 = style(doc, "Heading 1", size=26, before=24, after=30, page_break=True, outline=0, **head)
    style(doc, "Heading 2", size=15, before=20, after=8, outline=1, **head)
    h3 = style(doc, "Heading 3", size=12.5, before=14, after=6, outline=2, **head)
    h3.font.color.rgb = RGBColor.from_string(INK)
    style(doc, "SAUS Front Heading", size=26, before=24, after=30, page_break=True, outline=0, **head)
    style(doc, "SAUS Contents Heading", size=26, before=24, after=30, page_break=True, **head)
    style(doc, "SAUS Appendix Heading", size=26, before=24, after=30, page_break=True, outline=0,
          **head)
    for name in ("Heading 1", "SAUS Front Heading", "SAUS Contents Heading", "SAUS Appendix Heading"):
        st = styles[name]
        st.font.bold = False
        ppr = st.element.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:pBdr {W}><w:bottom w:val="single" w:sz="6" w:space="12" '
                             f'w:color="{RULE}"/></w:pBdr>'))
    for name in ("Heading 2", "Heading 3"):
        styles[name].font.bold = False
    numbering(doc)
    link_numbering(doc, "Heading 1", 80, 0)
    link_numbering(doc, "Heading 2", 80, 1)
    link_numbering(doc, "Heading 3", 80, 2)
    link_numbering(doc, "SAUS Appendix Heading", 81, 0)

    single = dict(line=1.0)
    style(doc, "SAUS Title", font=SEMIBOLD, size=24, color=MAROON, align=C, before=6, after=6, **single)
    style(doc, "SAUS Subtitle", font=LIGHT, size=15, color=SLATE, align=C, after=6, **single)
    style(doc, "SAUS Title Label", font=MEDIUM, size=10, color=GOLD, caps=True, spacing=30,
          align=C, before=14, after=4, keep_next=True, **single)
    style(doc, "SAUS Centred", align=C, after=0, **single)
    style(doc, "SAUS Header", font=SANS, size=9, color=SLATE, align=Lf, after=0, **single)
    styles["SAUS Header"].paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    style(doc, "SAUS Signature", font=SEMIBOLD, size=10.5, color=MAROON, align=Lf, before=40,
          after=0, keep_next=True, **single)
    style(doc, "SAUS Signature Detail", font=SANS, size=9.5, color=SLATE, align=Lf, after=10, **single)
    style(doc, "SAUS Dedication", size=14, italic=True, color=MAROON_DARK, align=C, before=200,
          page_break=True, **single)
    style(doc, "SAUS Keywords", align=Lf, before=10)
    style(doc, "SAUS Abbreviation", align=Lf, after=4, left=30, first=-30, **single)
    styles["SAUS Abbreviation"].paragraph_format.tab_stops.add_tab_stop(Mm(30))
    style(doc, "SAUS Reference", align=Lf, after=6, **single)
    link_numbering(doc, "SAUS Reference", 82)
    style(doc, "SAUS Code", font=MONO, size=9.5, align=Lf, after=0, left=4, keep_lines=True, **single)
    ppr = styles["SAUS Code"].element.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:pBdr {W}><w:left w:val="single" w:sz="18" w:space="6" '
                         f'w:color="{MAROON}"/></w:pBdr>'))
    ppr.append(parse_xml(f'<w:shd {W} w:val="clear" w:color="auto" w:fill="{MIST}"/>'))
    style(doc, "SAUS Figure", align=C, before=6, after=4, keep_next=True, **single)
    style(doc, "SAUS Equation", align=Lf, before=4, after=10, **single)
    eq = styles["SAUS Equation"].paragraph_format.tab_stops
    eq.add_tab_stop(Mm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    eq.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    style(doc, "SAUS Theorem", align=J, before=6, after=10)
    for name, font in (("SAUS Urdu", "SAUS Nastaliq Urdu"), ("SAUS Sindhi", "Lateef")):
        st = style(doc, name, size=14 if "Urdu" in name else 17, align=J, line=1.7)
        rp = st.element.get_or_add_rPr()
        rp.get_or_add_rFonts().set(qn("w:cs"), font)
        rp.append(parse_xml(f'<w:rtl {W}/>'))
        rp.append(parse_xml(f'<w:szCs {W} w:val="{28 if "Urdu" in name else 34}"/>'))
        st.element.get_or_add_pPr().append(parse_xml(f'<w:bidi {W}/>'))
    style(doc, "Caption", font=SANS, size=10, color=INK, align=C, before=4, after=14, bold=False,
          italic=False, **single)
    # contents: number, tab to the title, dotted tab to a right-aligned page number
    from docx.enum.text import WD_TAB_LEADER
    for name, fmt, left, hang in (
            ("toc 1", dict(font=SEMIBOLD, size=11, color=MAROON, before=9), 38, 38),
            ("toc 2", dict(size=12), 22, 12),
            ("toc 3", dict(size=12), 36, 14),
            ("toc 4", dict(font=SEMIBOLD, size=11, color=MAROON, before=9), 0, 0),
            ("table of figures", dict(size=12, after=4), 22, 22)):
        fmt.setdefault("after", 2)
        st = style(doc, name, align=Lf, left=left, first=-hang, **fmt, **single)
        tabs = st.paragraph_format.tab_stops
        if left:
            tabs.add_tab_stop(Mm(left))
        tabs.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    for name, num_id in (("List Bullet", 90), ("List Number", 91)):
        link_numbering(doc, name, num_id)
        st = styles[name]
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.alignment = Lf
        st.quick_style = True

    char = WD_STYLE_TYPE.CHARACTER
    style(doc, "SAUS Caption Label", kind=char, base=None, font=SEMIBOLD, color=MAROON)
    style(doc, "SAUS Key Term", kind=char, base=None, bold=True, color=MAROON)
    style(doc, "SAUS Code Inline", kind=char, base=None, font=MONO, size=10.5)
    style(doc, "SAUS Label", kind=char, base=None, font=SEMIBOLD, color=MAROON)
    style(doc, "SAUS Degree", kind=char, base=None, bold=True)
    style(doc, "Placeholder Text", kind=char, base=None, color="808080")

    styles.element.append(parse_xml(
        f'<w:style {W} w:type="table" w:customStyle="1" w:styleId="SAUSTable">'
        f'<w:name w:val="SAUS Table"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="59"/>'
        f'<w:qFormat/><w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'<w:jc w:val="left"/></w:pPr><w:rPr><w:sz w:val="22"/></w:rPr>'
        f'<w:tblPr><w:jc w:val="center"/><w:tblBorders>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="{INK}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{INK}"/></w:tblBorders>'
        f'<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        f'<w:bottom w:w="40" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar></w:tblPr>'
        f'<w:tblStylePr w:type="firstRow"><w:rPr><w:rFonts w:ascii="{SEMIBOLD}" w:hAnsi="{SEMIBOLD}"/>'
        f'</w:rPr><w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="4" w:space="0" '
        f'w:color="{INK}"/></w:tcBorders></w:tcPr></w:tblStylePr></w:style>'))
    L.setup_settings(doc)


# ---- Building blocks -----------------------------------------------------------------------------
def para(doc, text="", style_name="Normal", **fmt):
    p = doc.add_paragraph(style=style_name)
    if text:
        add_runs(p, (text, fmt))
    return p


def rich(doc, parts, style_name="Normal"):
    """parts: strings, or (text, fmt) tuples; fmt 'key' for key terms, 'code' for inline code."""
    p = doc.add_paragraph(style=style_name)
    for part in parts:
        text, fmt = (part, {}) if isinstance(part, str) else part
        if fmt == "key":
            fmt = dict(style="SAUSKeyTerm")
        elif fmt == "code":
            fmt = dict(style="SAUSCodeInline")
        add_runs(p, (text, fmt))
    return p


def fill(p, text, alias, tag, filled, **fmt):
    """Text in the demo; a grey fill-in field in the templates."""
    if filled:
        add_runs(p, (filled, fmt))
    else:
        # a no-break space before the field, so the gap survives next to the control
        texts = p._p.findall(f"{qn('w:r')}/{qn('w:t')}")
        if texts and texts[-1].text and texts[-1].text.endswith(" "):
            texts[-1].text = texts[-1].text[:-1] + " "
        sdt_inline(p, alias, tag, text, placeholder=True, **fmt)


def picture(doc, path, width_mm=None, height_mm=None, style_name="SAUS Figure"):
    p = doc.add_paragraph(style=style_name)
    kw = {"width": Mm(width_mm)} if width_mm else {"height": Mm(height_mm)}
    p.add_run().add_picture(path, **kw)
    return p


def caption(doc, kind, text):
    """'Figure 3.1  |  text' with Word's own numbering (chapter.number)."""
    p = doc.add_paragraph(style="Caption")
    fmt = dict(style="SAUSCaptionLabel")
    add_runs(p, (kind + " ", fmt))
    for r in field("STYLEREF 1 \\s", "1", **fmt):
        p._p.append(r)
    add_runs(p, (".", fmt))
    for r in field(f"SEQ {kind} \\* ARABIC \\s 1", "1", **fmt):
        p._p.append(r)
    add_runs(p, (" | " + text, {}))
    if kind == "Table":
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(6)
    return p


def toc(doc, instr, placeholder):
    p = doc.add_paragraph(style="Normal")
    for r in field(instr, placeholder):
        p._p.append(r)
    return p


def styleref(p, style_name, cached, **fmt):
    for r in field(f'STYLEREF "{style_name}"', cached, **fmt):
        p._p.append(r)


def rule_png(width_mm):
    path = os.path.join(HERE, f"rule-{width_mm}.png")
    dpmm = 600 / 25.4
    w, h = round(width_mm * dpmm), round(0.78 * dpmm)
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    m, g = round(5.67 * dpmm), round(2.48 * dpmm)
    d.rectangle([0, 0, m - 1, h - 1], fill="#" + MAROON)
    d.rectangle([m, 0, m + g - 1, h - 1], fill="#" + GOLD)
    hl = max(2, round(0.21 * dpmm))
    d.rectangle([m + g, (h - hl) // 2, w - 1, (h - hl) // 2 + hl - 1], fill="#" + RULE)
    img.save(path, dpi=(600, 600))
    return path


def signature_table(doc, entries):
    """Signature blocks, two to a row: entries = [(role, detail), ...]."""
    rows = (len(entries) + 1) // 2
    t = doc.add_table(rows=rows, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t._tbl.tblPr.append(parse_xml(f'<w:tblLayout {W} w:type="fixed"/>'))
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), (TEXT_W / 2, TEXT_W / 2)):
        col.set(qn("w:w"), str(int(w * 56.6929)))
    for i, (role, detail) in enumerate(entries):
        cell = t.cell(i // 2, i % 2)
        cell.width = Mm(TEXT_W / 2)
        tcpr = cell._tc.get_or_add_tcPr()
        tcpr.append(parse_xml(f'<w:tcMar {W}><w:left w:w="0" w:type="dxa"/>'
                              f'<w:right w:w="{int(14 * 56.69)}" w:type="dxa"/></w:tcMar>'))
        p = cell.paragraphs[0]
        p.style = "SAUS Signature"
        border(p, "top", color=INK, size=4, space=4)
        add_runs(p, (role, {}))
        d = cell.add_paragraph(style="SAUS Signature Detail")
        if isinstance(detail, tuple):             # a fill-in field
            sdt_inline(d, detail[0], detail[1], detail[2], placeholder=True)
        else:
            lines = detail.split("\n")
            for j, line in enumerate(lines):
                if j:
                    d.add_run().add_break(WD_BREAK.LINE)
                add_runs(d, (line, {}))
    return t


def equation(doc, number):
    """C = (L1 + 0.05) / (L2 + 0.05), centred, with its number on the right."""
    p = doc.add_paragraph(style="SAUS Equation")
    add_runs(p, ("\t", {}))

    def mr(t):
        return f"<m:r><m:t>{t}</m:t></m:r>"

    def lsub(i):
        return f"<m:sSub><m:e>{mr('L')}</m:e><m:sub>{mr(i)}</m:sub></m:sSub>"
    p._p.append(parse_xml(
        f'<m:oMath xmlns:m="{M}">{mr("C=")}<m:f><m:num>{lsub("1")}{mr("+0.05")}</m:num>'
        f'<m:den>{lsub("2")}{mr("+0.05")}</m:den></m:f></m:oMath>'))
    add_runs(p, ("\t(" + number + ")", {}))
    return p


def section_setup(sec):
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(38), Mm(25)
    sec.top_margin, sec.bottom_margin = Mm(30), Mm(27)
    sec.header_distance, sec.footer_distance = Mm(13), Mm(13)


def page_numbers(sec, fmt):
    sp = sec._sectPr
    for old in sp.findall(qn("w:pgNumType")):
        sp.remove(old)
    sp.find(qn("w:pgMar")).addnext(parse_xml(f'<w:pgNumType {W} w:fmt="{fmt}" w:start="1"/>'))


def clear_header_footer(sec):
    for hf in (sec.header, sec.footer):
        hf.is_linked_to_previous = False


def running_head(sec, part):
    """part: 'chapter' (Chapter 1 · Title), 'appendix' (Appendix A · Title) or a fixed text."""
    clear_header_footer(sec)
    hp = sec.header.paragraphs[0]
    hp.style = "SAUS Header"
    border(hp, "bottom", size=4, space=4)
    fmt = dict(size=9, color=SLATE)
    if part in ("chapter", "appendix"):
        ref = "1" if part == "chapter" else '"SAUS Appendix Heading"'
        for r in field(f"STYLEREF {ref} \\n", "Chapter 1" if part == "chapter" else "Appendix A",
                       **fmt):
            hp._p.append(r)
        add_runs(hp, (SEP, fmt))
        for r in field(f"STYLEREF {ref}", "Introduction", **fmt):
            hp._p.append(r)
    else:
        add_runs(hp, (part, fmt))
    add_runs(hp, ("\t", {}))
    for r in field("PAGE", "1", font=SEMIBOLD, color=MAROON, size=10):
        hp._p.append(r)


def new_part(doc, part):
    """A new section for the references or the appendices, numbered on from the chapters."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    section_setup(sec)
    sp = sec._sectPr
    for old in sp.findall(qn("w:pgNumType")):
        sp.remove(old)
    running_head(sec, part)
    return sec


# Order of child elements required by the WordprocessingML schema
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
           "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct "
           "topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind "
           "contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint "
           "noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u "
           "effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW "
             "jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
            "tcFitText vAlign hideMark",
}


def normalise(root):
    for tag, order in ORDER.items():
        rank = {qn("w:" + n): i for i, n in enumerate(order.split())}
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag, len(rank)))
            for k in kids:
                el.remove(k)
                el.append(k)


def normalise_all(doc):
    normalise(doc.element)
    normalise(doc.styles.element)
    normalise(doc.part.numbering_part.element)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            if not hf.is_linked_to_previous:
                normalise(hf._element)


# ---- The document ----------------------------------------------------------------------------------
def build(kind):
    """kind: 'fyp' (template), 'thesis' (template) or 'demo'."""
    demo = kind == "demo"
    fyp = kind in ("fyp", "demo")
    doctype = "project report" if fyp else "thesis"
    doc = Document()
    setup_styles(doc)
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)

    # -- Title page (no page number) -------------------------------------------------------
    sec = doc.sections[0]
    section_setup(sec)
    picture(doc, os.path.join(ASSETS, "saus-logo.png"), height_mm=32, style_name="SAUS Centred")
    p = para(doc, UNIVERSITY, "SAUS Centred", font=SEMIBOLD, size=14, color=MAROON, caps=True,
             spacing=30)
    p.paragraph_format.space_before = Pt(10)
    para(doc, MOTTO, "SAUS Centred", italic=True, color=MAROON_LIGHT)
    p = picture(doc, os.path.join(ART_PPT, "sindhi-maroon.png"), width_mm=48, style_name="SAUS Centred")
    p.paragraph_format.space_before = Pt(6)
    picture(doc, os.path.join(ART_PPT, "urdu-maroon.png"), width_mm=42, style_name="SAUS Centred")
    p = para(doc, "Final Year Project" if fyp else "Thesis", "SAUS Title Label")
    p.paragraph_format.space_before = Pt(26)
    p = doc.add_paragraph(style="SAUS Title")
    fill(p, "Title of the " + ("project" if fyp else "thesis"), "Title", "title",
         "A Shared Document Identity for The Shaikh Ayaz University" if demo else None)
    p = doc.add_paragraph(style="SAUS Subtitle")
    fill(p, "Subtitle (optional)", "Subtitle", "subtitle",
         "Templates for slides, posters, letters, certificates and theses" if demo else None)
    picture(doc, rule_png(50), width_mm=50, style_name="SAUS Centred").paragraph_format.space_before = Pt(10)
    p = para(doc, f"A {doctype} submitted in partial fulfilment of the requirements for the "
                  "degree of", "SAUS Centred")
    p.paragraph_format.space_before = Pt(24)
    p = doc.add_paragraph(style="SAUS Centred")
    fill(p, "Bachelor of Science in ..." if fyp else "Master of Science in ...", "Degree",
         "degree", "Bachelor of Science in Computer Science" if demo else None, style="SAUSDegree")
    para(doc, "Submitted by", "SAUS Title Label").paragraph_format.space_before = Pt(22)
    students = ([("Student One", "Roll No. 00-CS-001"), ("Student Two", "Roll No. 00-CS-002"),
                 ("Student Three", "Roll No. 00-CS-003")] if demo else
                [(None, None)] * (3 if fyp else 1))
    t = doc.add_table(rows=len(students), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (name, roll) in enumerate(students):
        for c, (val, ph, fmt) in enumerate(((name, "Student name", {}),
                                            (roll, "Roll No.", dict(color=SLATE)))):
            cp = t.cell(i, c).paragraphs[0]
            cp.style = "SAUS Centred"
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cp.paragraph_format.space_after = Pt(1)
            if c == 1:
                cp.paragraph_format.left_indent = Mm(8)
            fill(cp, ph, "Student", "student", val, **fmt)
    para(doc, "Supervised by", "SAUS Title Label").paragraph_format.space_before = Pt(14)
    p = doc.add_paragraph(style="SAUS Centred")
    fill(p, "Supervisor name", "Supervisor", "supervisor", "Dr. A. Supervisor" if demo else None)
    p = doc.add_paragraph(style="SAUS Centred")
    fill(p, "Designation, department", "Supervisor designation", "supdesig",
         "Assistant Professor, Department of Computer Science" if demo else None,
         size=10.5, color=SLATE)
    p = doc.add_paragraph(style="SAUS Centred")
    p.paragraph_format.space_before = Pt(24)
    border(p, "top", space=8)
    fill(p, "Department of ...", "Department", "department",
         "Department of Computer Science" if demo else None, size=11)
    para(doc, UNIVERSITY, "SAUS Centred", size=11)
    p = doc.add_paragraph(style="SAUS Centred")
    add_runs(p, ("Session ", dict(size=11, color=SLATE)))
    fill(p, "2022–2026", "Session", "session", "2022–2026" if demo else None,
         size=11, color=SLATE)
    add_runs(p, (SEP, dict(size=11, color=SLATE)))
    fill(p, "Month year", "Submission date", "submitted", "September 2026" if demo else None,
         size=11, color=SLATE)

    # -- Front matter (roman page numbers) ------------------------------------------------------
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    section_setup(sec)
    page_numbers(sec, "lowerRoman")
    clear_header_footer(sec)
    fp = sec.footer.paragraphs[0]
    fp.style = "SAUS Header"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in field("PAGE", "ii", font=SEMIBOLD, color=MAROON, size=10):
        fp._p.append(r)

    title_cached = "A Shared Document Identity for The Shaikh Ayaz University" if demo else "[title]"
    para(doc, "Certificate of Approval", "SAUS Front Heading")
    p = rich(doc, [f"This is to certify that the {doctype} entitled “"])
    styleref(p, "SAUS Title", title_cached, italic=True)
    add_runs(p, ("”, submitted by ", {}))
    fill(p, "names of the students" if fyp else "name of the student", "Students", "students",
         "Student One, Student Two and Student Three" if demo else None)
    add_runs(p, (", has been examined and approved as fulfilling the requirements for the degree "
                 "of ", {}))
    styleref(p, "SAUS Degree", "Bachelor of Science in Computer Science" if demo else "[degree]")
    add_runs(p, (".", {}))
    if demo:
        signature_table(doc, [("Supervisor", "Dr. A. Supervisor\nAssistant Professor, Department of Computer Science"),
                              ("External Examiner", "Name, designation and institution"),
                              ("Head of Department", "Department of Computer Science")])
    else:
        signature_table(doc, [("Supervisor", ("Supervisor", "sig1", "Name and designation")),
                              ("Co-supervisor (if any)", ("Co-supervisor", "sig2", "Name and designation")),
                              ("External Examiner", ("External examiner", "sig3", "Name and designation")),
                              ("Head of Department", ("Head of Department", "sig4", "Name and designation"))])

    para(doc, "Declaration", "SAUS Front Heading")
    we, our = ("We", "our") if fyp else ("I", "my")
    p = rich(doc, [f"{we} hereby declare that this {doctype}, entitled “"])
    styleref(p, "SAUS Title", title_cached, italic=True)
    add_runs(p, (f"”, is {our} own work, carried out under the supervision of ", {}))
    fill(p, "supervisor's name", "Supervisor", "supervisor2", "Dr. A. Supervisor" if demo else None)
    add_runs(p, (". It has not been submitted, in whole or in part, for any other degree or "
                 "qualification at this or any other institution. Work by others that is used in "
                 "it has been acknowledged and cited in the text and listed in the references.", {}))
    p = rich(doc, [f"The similarity index of this {doctype}, as reported by the university’s "
                   "plagiarism detection service, is "])
    fill(p, "..%", "Similarity index", "similarity", "9%" if demo else None)
    add_runs(p, (".", {}))
    signature_table(doc, [(n, r) for n, r in students] if demo else
                    [("Student name", ("Roll No.", "roll", "Roll No."))] * (3 if fyp else 1))

    p = doc.add_paragraph(style="SAUS Dedication")
    fill(p, "To ... (optional; delete this page if not needed)", "Dedication", "dedication",
         "To the teachers of Shikarpur." if demo else None)

    para(doc, "Acknowledgements", "SAUS Front Heading")
    if demo:
        para(doc, "We thank our supervisor for guidance throughout this project, the IT Directorate "
                  "for access to the university’s crest and photographs, and our classmates who "
                  "tested the templates in their own coursework.")
    else:
        sdt_block([para(doc, "Thank those who helped with the work.")], "Acknowledgements", "ack",
                  placeholder=True)

    para(doc, "Abstract", "SAUS Front Heading")
    if demo:
        para(doc, "Documents produced across a university—lecture slides, research posters, "
                  "official letters, certificates and theses—are often designed separately, and "
                  "so present the institution inconsistently. This project builds a single document "
                  "identity for The Shaikh Ayaz University, Shikarpur: a shared brand layer that holds "
                  "the colours sampled from the university crest, the typefaces and the institutional "
                  "details, and a family of LaTeX templates that draw on it. The templates set the "
                  "university’s name in English, Sindhi and Urdu, meet the WCAG 2.1 contrast "
                  "requirements for text, and compile with LuaLaTeX, XeLaTeX and pdfLaTeX. Matching "
                  "PowerPoint, Word and Excel templates extend the same identity to colleagues who do "
                  "not use LaTeX.")
    else:
        sdt_block([para(doc, "A one-page summary of the problem, the method, the results and the "
                             "conclusions.")], "Abstract", "abstract", placeholder=True)
    p = doc.add_paragraph(style="SAUS Keywords")
    add_runs(p, ("Keywords: ", dict(style="SAUSLabel")))
    fill(p, "keyword one, keyword two", "Keywords", "keywords",
         "document design, typography, LaTeX, accessibility, Sindhi, Urdu" if demo else None)

    para(doc, "Abstract in Urdu", "SAUS Front Heading")
    urdu = ("یہ منصوبہ شیخ ایاز یونیورسٹی، شکارپور کے لیے دستاویزات کے مشترکہ سانچوں کا ایک مجموعہ "
            "پیش کرتا ہے۔ ان سانچوں میں سلائیڈز، پوسٹر، سرکاری خطوط، اسناد اور مقالے شامل ہیں، جو "
            "یونیورسٹی کے نشان، رنگوں اور حروف کو یکساں انداز میں استعمال کرتے ہیں۔ یونیورسٹی کا نام "
            "انگریزی، سندھی اور اردو تینوں زبانوں میں درج کیا جاتا ہے۔")
    if demo:
        para(doc, urdu, "SAUS Urdu")
    else:
        sdt_block([para(doc, "خلاصہ یہاں لکھیں۔ (اگر ضرورت نہ ہو تو یہ صفحہ حذف کر دیں۔)",
                        "SAUS Urdu")], "Urdu abstract", "urdu", placeholder=True)

    para(doc, "Contents", "SAUS Contents Heading")
    # chapters from Heading 1-3; front matter and references at level 4 (no chapter
    # number column); appendices at level 1 like chapters
    toc(doc, 'TOC \\o "1-3" \\h \\z \\t "SAUS Front Heading,4,SAUS Appendix Heading,1"',
        "Right-click here and choose Update Field to build the contents.")
    para(doc, "List of Figures", "SAUS Front Heading")
    toc(doc, 'TOC \\h \\z \\c "Figure"', "Right-click here and choose Update Field.")
    para(doc, "List of Tables", "SAUS Front Heading")
    toc(doc, 'TOC \\h \\z \\c "Table"', "Right-click here and choose Update Field.")
    para(doc, "List of Abbreviations", "SAUS Front Heading")
    abbrs = ([("FYP", "Final Year Project"), ("HEC", "Higher Education Commission"),
              ("PDF", "Portable Document Format"), ("RTL", "Right to left"),
              ("SAUS", "The Shaikh Ayaz University, Shikarpur"),
              ("WCAG", "Web Content Accessibility Guidelines")] if demo else
             [("HEC", "Higher Education Commission"), ("SAUS", "The Shaikh Ayaz University, Shikarpur")])
    for a, meaning in abbrs:
        p = doc.add_paragraph(style="SAUS Abbreviation")
        add_runs(p, (a, dict(style="SAUSLabel")), ("\t" + meaning, {}))

    # -- Main matter (arabic page numbers, running head) ------------------------------------------
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    section_setup(sec)
    page_numbers(sec, "decimal")
    running_head(sec, "chapter")

    if demo:
        demo_chapters(doc)
    else:
        for chapter, sections in (("Introduction", ("Background", "Problem statement", "Objectives",
                                                    "Scope")),
                                  ("Literature Review", ()), ("Methodology", ()),
                                  ("Implementation", ()), ("Results and Discussion", ()),
                                  ("Conclusion and Future Work", ())):
            para(doc, chapter, "Heading 1")
            if not sections:
                para(doc, "Write this chapter here.")
            for s in sections:
                para(doc, s, "Heading 2")
                para(doc, "Write this section here.")
        new_part(doc, "References")
        para(doc, "References", "SAUS Front Heading")
        for ref in ("A. Author, Title of the Book, 2nd ed. City: Publisher, Year.",
                    "A. Author and B. Author, “Title of the article,” Journal, vol. 1, "
                    "no. 1, pp. 1–10, Year."):
            para(doc, ref, "SAUS Reference")
        new_part(doc, "appendix")
        para(doc, "Additional Material", "SAUS Appendix Heading")
        para(doc, "Appendices follow the references.")

    normalise_all(doc)
    props = doc.core_properties
    props.author = UNIVERSITY
    props.last_modified_by = UNIVERSITY   # not whoever last opened it in Word
    props.title = ("A Shared Document Identity for The Shaikh Ayaz University" if demo else
                   "The Shaikh Ayaz University, Shikarpur — " + doctype)
    props.subject = MOTTO
    props.comments = "Generated by word/source/build_thesis_docx.py"
    return doc


def demo_chapters(doc):
    """Chapters from latex/saus-thesis-demo.tex."""
    code = ("\\documentclass[aspectratio=169]{beamer}",
            "\\usetheme{SAUS}                  % colours, fonts and crest",
            "\\title{Digital Transformation of Campus Services}",
            "\\author{Dr. A. Presenter}",
            "\\begin{document}",
            "\\begin{frame}[plain]\\titlepage\\end{frame}",
            "\\end{document}")
    para(doc, "Introduction", "Heading 1")
    para(doc, "A university speaks through its documents. A student meets it first in an admission "
              "letter, a colleague elsewhere in a conference poster, and an employer in a degree "
              "certificate. When each of these is designed separately, the institution appears as "
              "several different organisations.")
    para(doc, "Problem statement", "Heading 2")
    para(doc, "The Shaikh Ayaz University, Shikarpur, a public sector general university with two "
              "faculties and seven degree programmes, had no shared templates for its documents. "
              "Departments prepared slides, posters and letters independently, with differing "
              "colours, typefaces and renderings of the university’s name.")
    para(doc, "Objectives", "Heading 2")
    para(doc, "The project set out to:")
    for item in ("derive a colour palette and typographic system from the university crest;",
                 "build templates for slides, posters, letters, certificates and theses that share "
                 "this system by construction, not by copying;",
                 "set the university’s name correctly in Sindhi and Urdu as well as English;",
                 "keep every text colour legible under WCAG 2.1 [1]."):
        para(doc, item, "List Number")
    para(doc, "Scope", "Heading 2")
    para(doc, "The templates target LaTeX [2], the standard for technical writing in the sciences, "
              "with matching Office templates for colleagues who work in PowerPoint, Word and Excel. "
              "Chapter 3 describes the design, Chapter 4 its implementation and Chapter 5 the "
              "evaluation.")

    para(doc, "Background", "Heading 1")
    para(doc, "Typography and identity", "Heading 2")
    para(doc, "Bringhurst describes typography as the craft of endowing language with a durable "
              "visual form [3]. For an institution, that form is part of its identity: the same "
              "typefaces, used the same way, make documents recognisably its own.")
    para(doc, "TeX and LaTeX", "Heading 2")
    para(doc, "TeX [4] separates the content of a document from its design. LaTeX classes and "
              "packages carry the design, so a single change to a class restyles every document "
              "built on it. KOMA-Script [5] provides the foundations for the letter and thesis "
              "classes, TikZ [6] draws the title pages, and fontspec [7] loads the typefaces.")
    para(doc, "Arabic-script names", "Heading 2")
    para(doc, "Sindhi and Urdu are written right to left in the Arabic script, Sindhi usually in "
              "Naskh and Urdu in Nastaliq. The templates set Sindhi in Lateef [8] and Urdu in Noto "
              "Nastaliq Urdu, both of which are bundled with the templates.")

    para(doc, "Design", "Heading 1")
    para(doc, "The crest", "Heading 2")
    para(doc, "The university crest (Figure 3.1) supplies the two primary colours: the maroon of "
              "its script and border, and the royal blue of its outer ring.")
    picture(doc, os.path.join(ASSETS, "saus-logo.png"), height_mm=45)
    caption(doc, "Figure", "The university crest, the source of the primary colours.")
    para(doc, "Colour and legibility", "Heading 2")
    p = rich(doc, [("Definition 3.1 (Contrast ratio).", dict(style="SAUSLabel")),
                   " For two colours with relative luminances L₁ ≥ L₂, the "
                   "contrast ratio is"], "SAUS Theorem")
    p.paragraph_format.space_after = Pt(2)
    equation(doc, "3.1")
    para(doc, "which runs from 1 (no contrast) to 21 (black on white) [1].")
    para(doc, "WCAG 2.1 asks for C ≥ 4.5 for body text and C ≥ 3 for large text. Table 3.1 "
              "lists the palette with each colour’s contrast against white, computed with "
              "Equation (3.1).")
    caption(doc, "Table", "The brand palette and its contrast against white.")
    table(doc, [("Colour", "Hex", "Contrast", "Use"),
                ("Maroon", "#6B2A4C", "10.2:1", "Primary"),
                ("Deep maroon", "#4A1C34", "13.9:1", "Large fields"),
                ("Royal blue", "#1A1A91", "13.1:1", "Secondary"),
                ("Gold", "#C39B3E", "2.6:1", "Rules and accents only"),
                ("Ink", "#231C21", "16.7:1", "Body text"),
                ("Slate", "#6A626B", "5.9:1", "Secondary text")],
          (30, 26, 22, 48), right={2}, mono={1})
    p = para(doc, "Gold falls below the threshold on white, so the templates use it only for rules "
                  "and ornaments, or for small capitals on the deep maroon field.")
    p.paragraph_format.space_before = Pt(10)

    para(doc, "Implementation", "Heading 1")
    para(doc, "A shared brand layer", "Heading 2")
    rich(doc, ["Every template loads one package, ", ("saus-brand.sty", "code"), ", which defines "
               "the colours, typefaces and identity (Figure 4.1). A change to that file reaches "
               "every template."])
    picture(doc, os.path.join(ART, "architecture.png"), width_mm=128)
    caption(doc, "Figure", "Every template draws on the shared brand layer.")
    para(doc, "Using a template", "Heading 2")
    para(doc, "A presentation needs only the theme; the university’s details are already built "
              "in. Listing 4.1 shows a complete title slide.")
    p = doc.add_paragraph(style="Caption")
    add_runs(p, ("Listing 4.1", dict(style="SAUSCaptionLabel")),
             (" | A title slide with the SAUS theme.", {}))
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    for line in code:
        para(doc, line, "SAUS Code")
    para(doc, "Engines", "Heading 2").paragraph_format.space_before = Pt(22)
    rich(doc, [("Proposition 4.1.", dict(style="SAUSLabel")),
               (" Each template compiles without errors under LuaLaTeX, XeLaTeX and pdfLaTeX.",
                dict(italic=True))], "SAUS Theorem")
    para(doc, "LuaLaTeX and XeLaTeX use the OpenType typefaces directly and set Sindhi and Urdu; "
              "pdfLaTeX falls back to Type 1 versions of the typefaces and omits the Arabic-script "
              "names, as Chapter 5 confirms.")

    para(doc, "Evaluation", "Heading 1")
    para(doc, "Compilation", "Heading 2")
    para(doc, "Each template’s demonstration file was compiled with all three engines. "
              "Table 5.1 summarises the result.")
    caption(doc, "Table", "Compilation of the demonstration files.")
    table(doc, [("Template", "LuaLaTeX", "XeLaTeX", "pdfLaTeX"),
                ("Slides", "Yes", "Yes", "Yesᵃ"), ("Posters", "Yes", "Yes", "Yesᵃ"),
                ("Letters", "Yes", "Yes", "Yesᵃ"), ("Certificates", "Yes", "Yes", "Yesᵃ")],
          (34, 24, 24, 24), centre={1, 2, 3})
    p = para(doc, "ᵃ Without the Sindhi and Urdu names.", "SAUS Centred", size=9.5, color=SLATE)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    para(doc, "Legibility", "Heading 2")
    para(doc, "Every colour used for text reaches a contrast of at least 4.5 against white "
              "(Table 3.1), satisfying Definition 3.1 at the level WCAG 2.1 requires for body text.")

    para(doc, "Conclusion", "Heading 1")
    para(doc, "The project produced a shared document identity for the university and a family of "
              "templates that apply it consistently. Because the design lives in one brand layer, "
              "future changes—a new typeface, a revised address—need to be made only once. "
              "Future work could extend the templates to the university’s website and to forms.")

    new_part(doc, "References")
    para(doc, "References", "SAUS Front Heading")
    refs = [
        ["World Wide Web Consortium, “Web content accessibility guidelines (WCAG) 2.1,” "
         "Jun. 5, 2018. [Online]. Available: https://www.w3.org/TR/WCAG21/"],
        ["L. Lamport, ", ("LaTeX: A Document Preparation System", True),
         ", 2nd ed. Reading, MA: Addison-Wesley, 1994."],
        ["R. Bringhurst, ", ("The Elements of Typographic Style", True),
         ", 4th ed. Vancouver: Hartley & Marks, 2012."],
        ["D. E. Knuth, ", ("The TeXbook", True), ". Reading, MA: Addison-Wesley, 1984."],
        ["M. Kohm, ", ("KOMA-Script: The Guide", True),
         ". [Online]. Available: https://ctan.org/pkg/koma-script"],
        ["T. Tantau, ", ("The TikZ and PGF Packages: Manual", True),
         ". [Online]. Available: https://ctan.org/pkg/pgf"],
        ["W. Robertson, ", ("The fontspec Package: Font Selection for XeLaTeX and LuaLaTeX", True),
         ". [Online]. Available: https://ctan.org/pkg/fontspec"],
        ["SIL Global, “Lateef.” [Online]. Available: https://software.sil.org/lateef/"]]
    for parts in refs:
        p = doc.add_paragraph(style="SAUS Reference")
        for part in parts:
            text, it = (part, False) if isinstance(part, str) else part
            add_runs(p, (text, dict(italic=it)))

    new_part(doc, "appendix")
    para(doc, "Using the templates", "SAUS Appendix Heading")
    rich(doc, ["The templates, with starter files and full documentation, are in the ",
               ("latex", "code"), " folder. Copy the template and its starter file, with ",
               ("saus-brand.sty", "code"), " and the ", ("assets", "code"),
               " folder, into a new directory and build with ", ("latexmk", "code"),
               ". The Word templates are in the ", ("word", "code"), " folder."])


def table(doc, rows, widths, right=(), centre=(), mono=()):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["SAUS Table"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 56.6929)))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Mm(widths[c])
            p = cell.paragraphs[0]
            p.style = "SAUS Centred"
            p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if c in right else
                           WD_ALIGN_PARAGRAPH.CENTER if c in centre else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            fmt = dict(size=11)
            if r == 0:
                fmt["font"] = SEMIBOLD
            elif c in mono:
                fmt["style"] = "SAUSCodeInline"
            add_runs(p, (value, fmt))
    return t


if __name__ == "__main__":
    L.save(build("fyp"), os.path.join(OUT, "SAUS-thesis-FYP.dotx"), True)
    L.save(build("thesis"), os.path.join(OUT, "SAUS-thesis.dotx"), True)
    L.save(build("demo"), os.path.join(OUT, "SAUS-thesis-demo.docx"), False)
    for f in os.listdir(HERE):
        if f.startswith("rule-") and f.endswith(".png"):
            os.remove(os.path.join(HERE, f))
    print("wrote SAUS-thesis-FYP.dotx, SAUS-thesis.dotx and SAUS-thesis-demo.docx")
